"""Compare live recommendation sensitivity across correct-score aggregators."""

from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wc_predictor.config import ProjectConfig
from wc_predictor.market_data import load_correct_score_odds, load_odds
from wc_predictor.paths import CACHE_CORRECT_SCORE_ODDS_PATH, OUTPUT_CORRECT_SCORE_AGGREGATION_XLSX_PATH
from wc_predictor.utils import ensure_parent_directory
from wc_predictor.workflow import run_prediction_workflow
from wc_predictor.world_cup import resolve_world_cup_odds_input

CORRECT_SCORE_AGGREGATION_METHODS = (
    "mean",
    "median",
    "trimmed_mean",
    "winsorized_mean",
    "reliability_weighted_mean",
    "auto",
)
DEFAULT_CORRECT_SCORE_ODDS_PATH = CACHE_CORRECT_SCORE_ODDS_PATH
DEFAULT_OUTPUT_PATH = OUTPUT_CORRECT_SCORE_AGGREGATION_XLSX_PATH

DETAIL_COLUMNS = [
    "match_id",
    "team_a",
    "team_b",
    "aggregation_method",
    "recommended_score",
    "best_expected_points",
    "top_5_ev_predictions",
    "top_10_market_scorelines",
    "top_10_blended_scorelines",
    "correct_score_bookmakers_count",
    "average_correct_score_overround",
    "max_correct_score_overround",
    "outlier_count",
    "kl_market_vs_poisson",
]


@dataclass(frozen=True)
class CorrectScoreAggregationComparison:
    """Detailed method runs and match-level recommendation sensitivity."""

    details: pd.DataFrame
    sensitivity: pd.DataFrame


def _build_sensitivity_summary(details: pd.DataFrame) -> pd.DataFrame:
    """Return one row per match plus one recommended-score column per method."""

    base = details[["match_id", "team_a", "team_b"]].drop_duplicates("match_id")
    pivot = details.pivot(index="match_id", columns="aggregation_method", values="recommended_score")
    pivot = pivot.reindex(columns=CORRECT_SCORE_AGGREGATION_METHODS).reset_index()
    pivot.columns.name = None
    grouped_scores = (
        details.groupby("match_id")["recommended_score"]
        .agg(lambda values: "; ".join(sorted(set(values.astype(str)))))
        .rename("unique_recommended_scores")
    )
    score_counts = (
        details.groupby("match_id")["recommended_score"]
        .nunique()
        .rename("number_of_unique_recommended_scores")
    )
    sensitivity = base.merge(pivot, on="match_id", validate="one_to_one")
    sensitivity = sensitivity.merge(grouped_scores, on="match_id", validate="one_to_one")
    sensitivity = sensitivity.merge(score_counts, on="match_id", validate="one_to_one")
    sensitivity["recommended_score_changes"] = sensitivity["number_of_unique_recommended_scores"] > 1
    sensitivity["aggregation_method_sensitive"] = sensitivity["recommended_score_changes"]
    return sensitivity.sort_values(["aggregation_method_sensitive", "match_id"], ascending=[False, True])


def compare_correct_score_aggregation_methods(
    odds_path: str | Path | None = None,
    correct_score_odds_path: str | Path = DEFAULT_CORRECT_SCORE_ODDS_PATH,
    *,
    poisson_weight: float = 0.85,
    config: ProjectConfig | None = None,
) -> CorrectScoreAggregationComparison:
    """Run the existing recommendation workflow under each aggregation method."""

    odds = load_odds(resolve_world_cup_odds_input(odds_path))
    correct_score_odds = load_correct_score_odds(correct_score_odds_path)
    base_config = config or ProjectConfig()
    rows: list[pd.DataFrame] = []
    for method in CORRECT_SCORE_AGGREGATION_METHODS:
        method_config = replace(
            base_config,
            correct_score_poisson_weight=poisson_weight,
            correct_score_aggregation_method=method,
        )
        report = run_prediction_workflow(odds, config=method_config, correct_score_odds=correct_score_odds).match_report
        report = report[report["has_correct_score_market"]].copy()
        report["aggregation_method"] = method
        report["kl_market_vs_poisson"] = report["correct_score_kl_divergence"]
        rows.append(report[DETAIL_COLUMNS])
    details = pd.concat(rows, ignore_index=True)
    details = details.sort_values(["match_id", "aggregation_method"], kind="stable").reset_index(drop=True)
    return CorrectScoreAggregationComparison(details, _build_sensitivity_summary(details))


def _style_worksheet(
    worksheet,
    *,
    decimal_columns: set[str] | None = None,
    wrapped_columns: set[str] | None = None,
) -> None:
    """Apply compact spreadsheet formatting to one comparison worksheet."""

    decimal_columns = decimal_columns or set()
    wrapped_columns = wrapped_columns or set()
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, cell in enumerate(worksheet[1], start=1):
        column_name = str(cell.value)
        for values in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for value in values:
                if column_name in decimal_columns:
                    value.number_format = "0.0000"
                if column_name in wrapped_columns:
                    value.alignment = Alignment(wrap_text=True, vertical="top")
        contents = [str(cell.value or "")]
        contents.extend(str(row[0].value or "") for row in worksheet.iter_cols(min_col=index, max_col=index, min_row=2))
        maximum_width = 60 if column_name in wrapped_columns else 32
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(max(len(content) for content in contents) + 2, 10),
            maximum_width,
        )


def export_correct_score_aggregation_comparison(
    comparison: CorrectScoreAggregationComparison,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> Path:
    """Write detailed and pivot-style sensitivity sheets to Excel."""

    output_path = Path(output_path)
    ensure_parent_directory(output_path)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        comparison.details.to_excel(writer, index=False, sheet_name="details")
        comparison.sensitivity.to_excel(writer, index=False, sheet_name="sensitivity")

    from openpyxl import load_workbook

    workbook = load_workbook(output_path)
    _style_worksheet(
        workbook["details"],
        decimal_columns={
            "best_expected_points",
            "average_correct_score_overround",
            "max_correct_score_overround",
            "kl_market_vs_poisson",
        },
        wrapped_columns={
            "top_5_ev_predictions",
            "top_10_market_scorelines",
            "top_10_blended_scorelines",
        },
    )
    _style_worksheet(
        workbook["sensitivity"],
        wrapped_columns={"unique_recommended_scores"},
    )
    workbook.save(output_path)
    return output_path


def format_correct_score_aggregation_comparison_summary(
    comparison: CorrectScoreAggregationComparison,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> str:
    """Render compact recommendation sensitivity diagnostics for the console."""

    sensitive = comparison.sensitivity[comparison.sensitivity["aggregation_method_sensitive"]]
    lines = [
        "Correct-Score Aggregation Comparison",
        f"- Methods compared: {', '.join(CORRECT_SCORE_AGGREGATION_METHODS)}",
        f"- Matches compared: {len(comparison.sensitivity)}",
        f"- Matches sensitive to aggregation method: {len(sensitive)}",
        "",
        "Match Sensitivity",
        "Match    Fixture                         Changes  Unique recommended scores",
    ]
    for _, row in comparison.sensitivity.iterrows():
        fixture = f"{row['team_a']} vs {row['team_b']}"
        lines.append(
            f"{row['match_id']:<8} {fixture:<31} "
            f"{'yes' if row['aggregation_method_sensitive'] else 'no':<7}  "
            f"{row['unique_recommended_scores']}"
        )
    lines.extend(["", f"Workbook written to {Path(output_path)}"])
    return "\n".join(lines)
