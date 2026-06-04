"""Report formatting and CSV/Excel exports."""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wc_predictor.backtesting import BatchBacktestReport, BatchBacktestSettings
from wc_predictor.utils import ensure_parent_directory


def _format_signed(value: float) -> str:
    return f"{value:+.3f}"


def _yes_no(value: bool) -> str:
    return "Yes" if value else "No"


def _strategy_average(summary: pd.DataFrame, strategy: str) -> float:
    row = summary.loc[summary["strategy"] == strategy]
    if row.empty:
        raise ValueError(f"Missing strategy summary: {strategy}")
    return float(row.iloc[0]["average_realised_points"])


def _format_overall_strategy_ranking(aggregate_summary: pd.DataFrame) -> str:
    if aggregate_summary.empty or not aggregate_summary["matches_used"].sum():
        return "Overall Strategy Ranking\nNo historical matches were processed."
    ranking = aggregate_summary.copy()
    favourite_average = _strategy_average(ranking, "favourite_1_0")
    modal_average = _strategy_average(ranking, "most_likely_poisson")
    ranking["gap_vs_favourite_1_0"] = ranking["average_realised_points"] - favourite_average
    ranking["gap_vs_most_likely_poisson"] = ranking["average_realised_points"] - modal_average
    ranking["rank"] = ranking["average_realised_points"].rank(method="min", ascending=False).astype(int)
    ranking = ranking.sort_values(
        ["average_realised_points", "total_points", "strategy"],
        ascending=[False, False, True],
    ).reset_index(drop=True)
    lines = [
        "Overall Strategy Ranking",
        "Rank  Strategy                       Avg pts  Total  Exact    Goal diff  Result   vs fav  vs modal",
    ]
    for _, row in ranking.iterrows():
        lines.append(
            f"{row['rank']:>4}  {row['strategy']:<29} "
            f"{row['average_realised_points']:>7.3f}  "
            f"{int(row['total_points']):>5}  "
            f"{row['exact_score_hit_rate']:>6.1%}  "
            f"{row['correct_goal_difference_hit_rate']:>9.1%}  "
            f"{row['correct_result_hit_rate']:>6.1%}  "
            f"{_format_signed(row['gap_vs_favourite_1_0']):>7}  "
            f"{_format_signed(row['gap_vs_most_likely_poisson']):>8}"
        )
    return "\n".join(lines)


def _format_key_conclusions(aggregate_summary: pd.DataFrame) -> str:
    if aggregate_summary.empty or not aggregate_summary["matches_used"].sum():
        return "Key Conclusions\n- No historical matches were processed."
    ranking = aggregate_summary.sort_values(
        ["average_realised_points", "total_points", "strategy"],
        ascending=[False, False, True],
    )
    best = ranking.iloc[0]
    tied_best = ranking[np.isclose(ranking["average_realised_points"], best["average_realised_points"])]
    best_strategies = ", ".join(sorted(tied_best["strategy"].astype(str)))
    favourite = _strategy_average(ranking, "favourite_1_0")
    modal = _strategy_average(ranking, "most_likely_poisson")
    ev_1x2 = _strategy_average(ranking, "ev_optimal_1x2")
    ev_over_under = _strategy_average(ranking, "ev_optimal_1x2_over_under")
    ev_beats_favourite = ev_1x2 > favourite
    ev_beats_modal = ev_1x2 > modal
    hypothesis_supported = ev_beats_favourite and ev_beats_modal
    return "\n".join(
        [
            "Key Conclusions",
            f"- Best strategy overall: {best_strategies} ({best['average_realised_points']:.3f} points per match)",
            f"- ev_optimal_1x2 beats favourite_1_0: {_yes_no(ev_beats_favourite)} ({_format_signed(ev_1x2 - favourite)} points per match)",
            f"- ev_optimal_1x2_over_under beats ev_optimal_1x2: {_yes_no(ev_over_under > ev_1x2)} ({_format_signed(ev_over_under - ev_1x2)} points per match)",
            f"- EV optimisation beats most_likely_poisson: {_yes_no(ev_beats_modal)} ({_format_signed(ev_1x2 - modal)} points per match)",
            f"- Results support the project hypothesis: {_yes_no(hypothesis_supported)}",
        ]
    )


def _per_file_winners(detailed_summary: pd.DataFrame) -> tuple[pd.DataFrame, pd.Series]:
    winners: list[dict[str, object]] = []
    first_place_counts: dict[str, int] = {}
    for source_file, group in detailed_summary.groupby("source_file", sort=True):
        ranked = group.sort_values(
            ["average_realised_points", "total_points", "strategy"],
            ascending=[False, False, True],
        ).reset_index(drop=True)
        best = ranked.iloc[0]
        tied_best = ranked[np.isclose(ranked["average_realised_points"], best["average_realised_points"])]
        tied_names = sorted(tied_best["strategy"].astype(str))
        for strategy in tied_names:
            first_place_counts[strategy] = first_place_counts.get(strategy, 0) + 1
        remaining = ranked[~np.isclose(ranked["average_realised_points"], best["average_realised_points"])]
        second = remaining.iloc[0] if not remaining.empty else None
        winners.append(
            {
                "source_file": source_file,
                "best_strategy": ", ".join(tied_names),
                "best_average_points": float(best["average_realised_points"]),
                "second_best_strategy": second["strategy"] if second is not None else "None",
                "performance_gap": (
                    float(best["average_realised_points"] - second["average_realised_points"])
                    if second is not None
                    else np.nan
                ),
            }
        )
    winners_frame = pd.DataFrame(winners)
    win_counts = pd.Series(first_place_counts, dtype=int).sort_index()
    return winners_frame, win_counts


def _format_per_file_winners(detailed_summary: pd.DataFrame) -> str:
    if detailed_summary.empty:
        return "Per-File Winners\nNo Football-Data CSV files were processed."
    winners, win_counts = _per_file_winners(detailed_summary)
    lines = [
        "Per-File Winners",
        "Source file                    Best strategy                   Avg pts  Second-best strategy            Gap",
    ]
    for _, row in winners.iterrows():
        lines.append(
            f"{row['source_file']:<30} "
            f"{row['best_strategy']:<29} "
            f"{row['best_average_points']:>7.3f}  "
            f"{row['second_best_strategy']:<29} "
            f"{_format_signed(row['performance_gap']) if pd.notna(row['performance_gap']) else 'n/a':>7}"
        )
    lines.append("")
    lines.append("First-Place Counts")
    for strategy, count in win_counts.items():
        lines.append(f"- {strategy}: {count}")
    return "\n".join(lines)


def _format_skipped_matches(skipped_by_file_reason: pd.DataFrame) -> str:
    if skipped_by_file_reason.empty:
        return "Skipped Matches\nNo skipped matches."
    compact = (
        skipped_by_file_reason.groupby("reason", as_index=False)
        .agg(
            skipped_matches=("skipped_matches", "sum"),
            files_affected=("source_file", "nunique"),
        )
        .sort_values(["skipped_matches", "reason"], ascending=[False, True])
    )
    lines = ["Skipped Matches", "Reason                               Matches  Files"]
    for _, row in compact.iterrows():
        lines.append(f"{row['reason']:<36} {int(row['skipped_matches']):>7}  {int(row['files_affected']):>5}")
    return "\n".join(lines)


def _format_favourite_strength_analysis(favourite_strength_summary: pd.DataFrame) -> str:
    lines = [
        "Favourite-Strength Analysis",
        "---------------------------",
        "Bucket              Matches  Best strategy                   Avg pts  EV vs fav_1_0",
    ]
    for _, row in favourite_strength_summary.iterrows():
        if int(row["matches"]) == 0:
            lines.append(f"{row['favourite_bucket']:<19} {0:>7}  {'n/a':<29} {'n/a':>7}  {'n/a':>13}")
            continue
        lines.append(
            f"{row['favourite_bucket']:<19} "
            f"{int(row['matches']):>7}  "
            f"{row['best_strategy']:<29} "
            f"{row['best_average_points']:>7.3f}  "
            f"{_format_signed(row['best_ev_gap_vs_favourite_1_0']):>13}"
        )
    return "\n".join(lines)


def format_backtest_console_summary(
    report: BatchBacktestReport,
    input_path: str | Path,
    settings: BatchBacktestSettings,
    verbose: bool = False,
) -> str:
    """Render the default human-readable historical-backtest console report."""

    source_files = report.detailed_summary["source_file"].nunique() if "source_file" in report.detailed_summary else 0
    skipped_input_files = len(report.skipped_files)
    total_matches = (
        report.predictions[["source_file", "match_id"]].drop_duplicates().shape[0]
        if not report.predictions.empty
        else 0
    )
    total_skipped = (
        report.skipped[["source_file", "match_id"]].drop_duplicates().shape[0]
        if not report.skipped.empty
        else 0
    )
    sections = [
        "\n".join(
            [
                "Backtest Scope",
                f"- Source: {Path(input_path)}",
                f"- Files processed: {source_files}",
                f"- Non-Football-Data CSV files skipped: {skipped_input_files}",
                f"- Total matches used: {total_matches}",
                f"- Total skipped matches: {total_skipped}",
            ]
        ),
        _format_overall_strategy_ranking(report.aggregate_summary),
        _format_key_conclusions(report.aggregate_summary),
        _format_favourite_strength_analysis(report.favourite_strength_summary),
        _format_per_file_winners(report.detailed_summary),
        _format_skipped_matches(report.skipped_by_file_reason),
        "\n".join(
            [
                "File Outputs",
                f"- Detailed per-file CSV: {settings.detailed_output_path}",
                f"- Aggregate strategy CSV: {settings.aggregate_output_path}",
                f"- Skipped-match CSV: {settings.skipped_output_path}",
                f"- Favourite-strength CSV: {settings.favourite_strength_output_path}",
            ]
        ),
    ]
    if verbose:
        sections.extend(
            [
                "Verbose Per-File Strategy Results\n" + report.detailed_summary.to_string(index=False),
                "Verbose Aggregate Strategy Results\n" + report.aggregate_summary.to_string(index=False),
                "Verbose Skipped Matches By File And Reason\n"
                + (
                    report.skipped_by_file_reason.to_string(index=False)
                    if not report.skipped_by_file_reason.empty
                    else "No skipped matches."
                ),
                "Verbose Favourite-Strength Results\n"
                + report.favourite_strength_summary.to_string(index=False),
            ]
        )
    return "\n\n".join(sections)


def format_model_inspection_report(match_report: pd.DataFrame) -> str:
    """Render a compact manual-inspection report from workflow match output."""

    sections: list[str] = []
    for _, row in match_report.iterrows():
        qualifier = f"; qualifier={row['recommended_qualifier']}" if row.get("recommended_qualifier") else ""
        warnings = f"\n  Warnings: {row['warnings']}" if row.get("warnings") else ""
        correct_score_diagnostics = (
            [
                f"  Correct-score Poisson weight: {row['correct_score_poisson_weight']:.2f}",
                f"  Correct-score blend applied: {'yes' if row['correct_score_blend_applied'] else 'no'}",
                f"  Correct-score blend note: {row['correct_score_blend_note'] or 'none'}",
                f"  Correct-score coverage: {row['correct_score_scorelines_count']} scorelines / {row['correct_score_bookmakers_count']} bookmakers",
                f"  Correct-score sparse warning: {row['correct_score_sparse_warning'] or 'none'}",
                f"  Correct-score aggregation: {row['correct_score_aggregation_method']} ({row['outlier_count']} outliers)",
                f"  Correct-score overround: avg={row['average_correct_score_overround']:.4f} max={row['max_correct_score_overround']:.4f}",
                f"  Correct-score bookmakers: {row['correct_score_bookmaker_diagnostics']}",
                f"  Correct-score coverage warning: {row['scoreline_coverage_warning'] or 'none'}",
                f"  Correct-score out-of-grid scorelines: {row['out_of_grid_scorelines_count']}",
                f"  Correct-score top outliers: {row['top_outlier_examples'] or 'none'}",
                f"  Correct-score market top 10: {row['correct_score_market_top_10']}",
                f"  Blended score top 10: {row['correct_score_blended_top_10']}",
                f"  Correct-score market-to-Poisson KL divergence: {row['correct_score_kl_divergence']:.6f}",
            ]
            if row.get("has_correct_score_market")
            else []
        )
        sections.append(
            "\n".join(
                [
                    f"{row['match_id']} | {row['stage']} | {row['team_a']} vs {row['team_b']}",
                    f"  Raw 1X2 by bookmaker: {row['bookmaker_raw_1x2']}",
                    f"  Fair 1X2 by bookmaker: {row['bookmaker_fair_1x2']}",
                    f"  Aggregated fair 1X2: A={row['market_a_win']:.4f} D={row['market_draw']:.4f} B={row['market_b_win']:.4f}",
                    f"  Favourite strength: p_fav={row['favourite_probability']:.4f} ({row['favourite_bucket']})",
                    f"  Calibrated lambdas: A={row['lambda_a']:.4f} B={row['lambda_b']:.4f}",
                    f"  Model-implied 1X2: A={row['model_a_win']:.4f} D={row['model_draw']:.4f} B={row['model_b_win']:.4f}",
                    f"  Calibration error: {row['calibration_loss']:.6f}",
                    f"  Total-goals lines available: {row['total_goals_lines_available'] or 'none'}",
                    f"  Total-goals lines used for calibration: {row['total_goals_lines_used_for_calibration'] or 'none'}",
                    f"  Total-goals lines skipped for calibration: {row['total_goals_lines_skipped_for_calibration'] or 'none'}",
                    f"  Total-goals line fit error: {row['total_goals_line_fit_error'] if pd.notna(row['total_goals_line_fit_error']) else 'n/a'}",
                    f"  Total-goals line diagnostics: {row['total_goals_line_diagnostics'] or 'none'}",
                    f"  O/U 2.5 used: {'yes' if row['over_under_2_5_used'] else 'no'}",
                    f"  Multi-line totals used: {'yes' if row['multi_line_totals_used'] else 'no'}",
                    f"  Score-matrix tail mass: {row['tail_probability_before_renormalisation']:.6%}",
                    f"  Most likely scoreline: {row['most_likely_scoreline']}",
                    f"  EV-optimal prediction: {row['recommended_score']}{qualifier} ({row['best_expected_points']:.3f} EV)",
                    f"  Top 5 EV predictions: {row['top_5_ev_predictions']}",
                    f"  Model comparison: baseline={row['baseline_poisson_recommended_score']} "
                    f"correct_score_blended={row['correct_score_blended_recommended_score']} "
                    f"final_live={row['final_live_recommended_score']} "
                    f"agree={'yes' if row['model_recommendations_agree'] else 'no'}",
                    f"  Dixon-Coles challenger: rho={row['dixon_coles_rho']:.4f} "
                    f"recommended={row['dixon_coles_recommended_score']} "
                    f"changes_final={'yes' if row['dixon_coles_changes_recommendation'] else 'no'}",
                    f"  Dixon-Coles top 5 EV predictions: {row['dixon_coles_top_5_ev_predictions']}",
                    *correct_score_diagnostics,
                ]
            )
            + warnings
        )
    return "\n\n".join(sections)


def format_world_cup_console_summary(
    match_report: pd.DataFrame,
    input_path: str | Path,
    csv_output_path: str | Path,
    xlsx_output_path: str | Path,
    submission_xlsx_output_path: str | Path,
) -> str:
    """Render the real-tournament recommendation summary."""

    sections = [
        "\n".join(
            [
                "World Cup Predictions",
                f"- Source: {Path(input_path)}",
                f"- Matches: {len(match_report)}",
            ]
        )
    ]
    for _, row in match_report.iterrows():
        group = f" / {row['group']}" if pd.notna(row.get("group")) and str(row["group"]).strip() else ""
        qualifier = f"; qualifier={row['recommended_qualifier']}" if row.get("recommended_qualifier") else ""
        correct_score_diagnostics = (
            [
                f"  Correct-score blend: poisson_weight={row['correct_score_poisson_weight']:.2f} KL={row['correct_score_kl_divergence']:.6f}",
                f"  Correct-score blend applied: {'yes' if row['correct_score_blend_applied'] else 'no'}",
                f"  Correct-score blend note: {row['correct_score_blend_note'] or 'none'}",
                f"  Correct-score coverage: {row['correct_score_scorelines_count']} scorelines / {row['correct_score_bookmakers_count']} bookmakers",
                f"  Correct-score sparse warning: {row['correct_score_sparse_warning'] or 'none'}",
                f"  Correct-score aggregation: {row['correct_score_aggregation_method']} ({row['outlier_count']} outliers)",
                f"  Correct-score overround: avg={row['average_correct_score_overround']:.4f} max={row['max_correct_score_overround']:.4f}",
                f"  Correct-score bookmakers: {row['correct_score_bookmaker_diagnostics']}",
                f"  Correct-score coverage warning: {row['scoreline_coverage_warning'] or 'none'}",
                f"  Correct-score out-of-grid scorelines: {row['out_of_grid_scorelines_count']}",
                f"  Correct-score top outliers: {row['top_outlier_examples'] or 'none'}",
                f"  Correct-score market top 10: {row['correct_score_market_top_10']}",
                f"  Blended score top 10: {row['correct_score_blended_top_10']}",
            ]
            if row.get("has_correct_score_market")
            else []
        )
        sections.append(
            "\n".join(
                [
                    f"{row['match_id']} | {row['team_a']} vs {row['team_b']}",
                    f"  Stage/group: {row['stage']}{group}",
                    f"  Fair 1X2: {row['team_a']}={row['market_a_win']:.2%} Draw={row['market_draw']:.2%} {row['team_b']}={row['market_b_win']:.2%}",
                    f"  Calibrated lambdas: {row['team_a']}={row['lambda_a']:.4f} {row['team_b']}={row['lambda_b']:.4f}",
                    f"  Total-goals lines: available={row['total_goals_lines_available'] or 'none'} used={row['total_goals_lines_used_for_calibration'] or 'none'}",
                    f"  Total-goals fit error: {row['total_goals_line_fit_error'] if pd.notna(row['total_goals_line_fit_error']) else 'n/a'}; O/U 2.5 used={'yes' if row['over_under_2_5_used'] else 'no'}; multi-line used={'yes' if row['multi_line_totals_used'] else 'no'}",
                    f"  Total-goals line diagnostics: {row['total_goals_line_diagnostics'] or 'none'}",
                    f"  Favourite strength: p_fav={row['favourite_probability']:.2%} ({row['favourite_bucket']})",
                    f"  Modal scoreline: {row['most_likely_scoreline']}",
                    f"  EV-optimal scoreline: {row['recommended_score']}{qualifier}",
                    f"  Top 5 EV scorelines: {row['top_5_ev_predictions']}",
                    f"  Model comparison: baseline={row['baseline_poisson_recommended_score']} "
                    f"correct_score_blended={row['correct_score_blended_recommended_score']} "
                    f"final_live={row['final_live_recommended_score']} "
                    f"agree={'yes' if row['model_recommendations_agree'] else 'no'}",
                    f"  Dixon-Coles challenger: rho={row['dixon_coles_rho']:.4f} "
                    f"recommended={row['dixon_coles_recommended_score']} "
                    f"changes_final={'yes' if row['dixon_coles_changes_recommendation'] else 'no'}",
                    f"  Dixon-Coles top 5 EV scorelines: {row['dixon_coles_top_5_ev_predictions']}",
                    f"  Public-ranking strategy: mode={row['public_strategy_mode']} "
                    f"score={row['public_strategy_score']} "
                    f"EV_cost={row['public_strategy_ev_cost']:.3f}",
                    f"  Estimated most crowded public score: {row['estimated_most_crowded_public_score']} "
                    f"({row['estimated_most_crowded_public_pick_share']:.2%})",
                    f"  Public strategy score: {row['public_strategy_score']} "
                    f"public_share={row['public_strategy_public_pick_share']:.2%} "
                    f"leverage={row['public_strategy_leverage_score']:.3f}",
                    *correct_score_diagnostics,
                ]
            )
        )
    sections.append(
        "\n".join(
            [
                "File Outputs",
                f"- CSV recommendations: {Path(csv_output_path)}",
                f"- Excel recommendations: {Path(xlsx_output_path)}",
                f"- Submission sheet: {Path(submission_xlsx_output_path)}",
            ]
        )
    )
    sections.append(format_world_cup_prediction_diagnostics(match_report))
    sections.append(format_world_cup_model_risk_summary(match_report))
    return "\n\n".join(sections)


def format_world_cup_prediction_diagnostics(match_report: pd.DataFrame) -> str:
    """Render compact post-generation diagnostics for real-tournament output."""

    score_counts = match_report["recommended_score"].value_counts().sort_index()
    bucket_counts = match_report["favourite_bucket"].value_counts().sort_index()
    differs_count = int(match_report["ev_optimal_differs_from_most_likely"].sum())
    lines = [
        "Prediction Output Diagnostics",
        "Recommended Score Counts",
        *(f"- {score}: {count}" for score, count in score_counts.items()),
        "",
        "Favourite-Strength Bucket Counts",
        *(f"- {bucket}: {count}" for bucket, count in bucket_counts.items()),
        "",
        f"- Average lambda_a: {match_report['lambda_a'].mean():.4f}",
        f"- Average lambda_b: {match_report['lambda_b'].mean():.4f}",
        f"- EV-optimal score differs from modal scoreline: {differs_count}",
        "",
        "Top 10 Matches By Favourite Probability",
    ]
    strongest_favourites = match_report.nlargest(10, "favourite_probability")
    for _, row in strongest_favourites.iterrows():
        lines.extend(
            [
                f"{row['match_id']} | {row['team_a']} vs {row['team_b']} | p_fav={row['favourite_probability']:.2%}",
                f"  Fair 1X2: market_a={row['market_a_win']:.2%} market_draw={row['market_draw']:.2%} market_b={row['market_b_win']:.2%}",
                f"  Lambdas: lambda_a={row['lambda_a']:.4f} lambda_b={row['lambda_b']:.4f}",
                f"  Scorelines: recommended={row['recommended_score']} modal={row['most_likely_scoreline']}",
                f"  Top 5 EV predictions: {row['top_5_ev_predictions']}",
            ]
        )
    return "\n".join(lines)


def format_world_cup_model_risk_summary(match_report: pd.DataFrame) -> str:
    """Render compact data-quality and model-risk counts for live review."""

    flags = match_report["warning_flags"].fillna("").map(lambda value: set(str(value).split("; ")))
    count_flag = lambda flag: int(flags.map(lambda values: flag in values).sum())
    return "\n".join(
        [
            "Model-Risk Summary",
            f"- Matches: {len(match_report)}",
            f"- Matches with only one bookmaker: {count_flag('only_one_bookmaker')}",
            f"- Matches without over/under odds: {count_flag('no_over_under')}",
            f"- Matches without BTTS odds: {count_flag('no_btts')}",
            f"- Extreme favourites: {count_flag('extreme_favourite')}",
            f"- EV-optimal score differs from modal scoreline: {int(match_report['ev_optimal_differs_from_most_likely'].sum())}",
            f"- Matches with high calibration error: {count_flag('high_calibration_error')}",
            f"- Matches with lambda_a near a calibration bound: {count_flag('lambda_a_near_bound')}",
            f"- Matches with lambda_b near a calibration bound: {count_flag('lambda_b_near_bound')}",
            f"- Sparse correct-score markets with blending suppressed: {count_flag('correct_score_blend_suppressed_sparse_market')}",
            f"- Matches with challenger-model recommendation disagreement: {count_flag('challenger_model_recommendations_disagree')}",
            f"- Knockout matches missing qualification odds: {count_flag('knockout_missing_qualification_odds')}",
        ]
    )


def export_dataframe(frame: pd.DataFrame, path: str | Path) -> None:
    """Export a report as CSV or Excel according to the file suffix."""

    path = Path(path)
    ensure_parent_directory(path)
    if path.suffix.lower() == ".csv":
        frame.to_csv(path, index=False)
    elif path.suffix.lower() == ".xlsx":
        frame.to_excel(path, index=False)
    else:
        raise ValueError(f"Unsupported report file type: {path.suffix}")


def _ev_decomposition_sheet(frame: pd.DataFrame) -> pd.DataFrame:
    source_column = "top_10_ev_decomposition_json" if "top_10_ev_decomposition_json" in frame else "top_5_ev_decomposition_json"
    if source_column not in frame:
        return pd.DataFrame()
    records: list[dict[str, object]] = []
    for _, row in frame.iterrows():
        raw_records = json.loads(str(row[source_column] or "[]"))
        for record in raw_records:
            record = dict(record)
            record.setdefault("match_id", row["match_id"])
            record.setdefault("team_a", row["team_a"])
            record.setdefault("team_b", row["team_b"])
            records.append(record)
    return pd.DataFrame(records)


def _high_score_diagnostics_sheet(frame: pd.DataFrame) -> pd.DataFrame:
    if "extreme_favourite_audit_triggered" not in frame:
        return pd.DataFrame()
    columns = [
        "match_id",
        "team_a",
        "team_b",
        "recommended_score",
        "favourite_probability",
        "current_final_recommendation",
        "normal_grid_poisson_recommendation",
        "larger_grid_poisson_recommendation",
        "normal_grid_recommendation",
        "larger_grid_recommendation",
        "recommendation_changes_with_larger_grid",
        "larger_grid_poisson_changes_recommendation",
        "larger_grid_differs_from_live_recommendation",
        "normal_grid_tail_mass",
        "larger_grid_tail_mass",
        "ev_favourite_3_0_normal_grid",
        "ev_favourite_4_0_normal_grid",
        "ev_favourite_5_0_normal_grid",
        "ev_favourite_3_0_larger_grid",
        "ev_favourite_4_0_larger_grid",
        "ev_favourite_5_0_larger_grid",
        "top_clean_sheet_scores",
        "top_favourite_margin_probabilities",
        "top_5_favourite_score_count_probabilities",
        "high_score_cluster_scores",
        "three_four_five_nil_within_0_10_ev",
        "warning_flags",
    ]
    diagnostics = frame[frame["extreme_favourite_audit_triggered"]].copy()
    return diagnostics[[column for column in columns if column in diagnostics]]


def _market_consistent_diagnostics_sheet(frame: pd.DataFrame) -> pd.DataFrame:
    if "market_consistent_recommended_score" not in frame:
        return pd.DataFrame()
    columns = [
        "match_id",
        "team_a",
        "team_b",
        "recommended_score",
        "market_consistent_recommended_score",
        "market_consistent_differs_from_default",
        "market_consistent_status",
        "market_consistent_optimisation_classification",
        "market_consistent_optimisation_success",
        "market_consistent_optimisation_status_code",
        "market_consistent_optimisation_message",
        "market_consistent_optimisation_iterations",
        "market_consistent_final_objective_value",
        "market_consistent_gradient_norm",
        "market_consistent_max_constraint_error",
        "market_consistent_constraint_count",
        "market_consistent_1x2_constraint_count",
        "market_consistent_btts_constraint_count",
        "market_consistent_total_goals_constraint_count",
        "market_consistent_asian_handicap_constraint_count",
        "market_consistent_correct_score_constraint_count",
        "warning_flags",
        "market_consistent_best_expected_points",
        "market_consistent_ev_gap_best_vs_second",
        "market_consistent_kl_divergence_vs_prior",
        "market_consistent_1x2_fit_error",
        "market_consistent_btts_fit_error",
        "market_consistent_total_goals_fit_error",
        "market_consistent_asian_handicap_fit_error",
        "market_consistent_asian_handicap_fit_error_selected",
        "market_consistent_asian_handicap_fit_error_all",
        "market_consistent_largest_margin_shift_value",
        "market_consistent_correct_score_fit_error",
        "market_consistent_asian_totals_used",
        "market_consistent_asian_totals_shift_recommendation",
        "market_consistent_asian_handicap_lines_available",
        "market_consistent_asian_handicap_lines_selected",
        "market_consistent_asian_handicap_lines_skipped",
        "market_consistent_asian_handicap_lines_used",
        "market_consistent_asian_handicap_lines_skipped_detail",
        "asian_handicap_shift_recommendation",
        "market_consistent_margin_distribution_before",
        "market_consistent_margin_distribution_after",
        "market_consistent_largest_margin_shift",
        "market_consistent_top_10_ev_scorelines",
        "market_consistent_top_10_ev_decomposition",
        "market_consistent_top_10_probability_scorelines",
        "market_consistent_poisson_ev_favourite_3_0",
        "market_consistent_poisson_ev_favourite_4_0",
        "market_consistent_poisson_ev_favourite_5_0",
        "market_consistent_matrix_ev_favourite_3_0",
        "market_consistent_matrix_ev_favourite_4_0",
        "market_consistent_matrix_ev_favourite_5_0",
    ]
    selected_columns = list(dict.fromkeys(column for column in columns if column in frame))
    return frame[selected_columns].copy()


def export_world_cup_recommendations_excel(
    frame: pd.DataFrame,
    path: str | Path,
    margin_method_comparison: pd.DataFrame | None = None,
    final_decision_dashboard: pd.DataFrame | None = None,
    asian_handicap_diagnostics: pd.DataFrame | None = None,
    margin_diagnostics: pd.DataFrame | None = None,
) -> None:
    """Export readable real-tournament recommendations with Excel formatting."""

    path = Path(path)
    ensure_parent_directory(path)
    aliases = {
        "market_a_win": "market_a",
        "market_b_win": "market_b",
        "ev_optimal_differs_from_most_likely": "ev_optimal_differs",
    }
    key_columns = [
        "match_id",
        "date",
        "time",
        "stage",
        "group",
        "team_a",
        "team_b",
        "recommended_score",
        "final_decision_score",
        "confidence_level",
        "manual_review_flag",
        "main_alternative_score",
        "plausible_alternatives",
        "strategic_alternative_score",
        "override_candidate",
        "model_consensus",
        "risk_notes",
        "decision_note",
        "market_a_win",
        "market_draw",
        "market_b_win",
        "favourite_probability",
        "favourite_bucket",
        "lambda_a",
        "lambda_b",
        "recommended_qualifier",
        "best_expected_points",
        "most_likely_scoreline",
        "ev_optimal_differs_from_most_likely",
        "top_5_ev_predictions",
        "plausible_top_alternatives",
        "suppressed_ev_candidates",
        "suppression_reason",
        "recommendation_confidence",
        "manual_review_flag",
        "close_alternatives",
        "ev_gap_to_second",
        "ev_gap_to_third",
        "decision_note",
        "expected_team_a_goals",
        "expected_team_b_goals",
        "expected_total_goals",
        "probability_team_a_scores_0",
        "probability_team_a_scores_1",
        "probability_team_a_scores_2",
        "probability_team_a_scores_3",
        "probability_team_a_scores_4",
        "probability_team_a_scores_5",
        "probability_team_a_scores_6_plus",
        "probability_team_b_scores_0",
        "probability_team_b_scores_1",
        "probability_team_b_scores_2",
        "probability_team_b_scores_3",
        "probability_team_b_scores_4",
        "probability_team_b_scores_5",
        "probability_team_b_scores_6_plus",
        "probability_total_goals_0",
        "probability_total_goals_1",
        "probability_total_goals_2",
        "probability_total_goals_3",
        "probability_total_goals_4",
        "probability_total_goals_5",
        "probability_total_goals_6",
        "probability_total_goals_7_plus",
        "probability_team_a_wins_by_1",
        "probability_team_a_wins_by_2",
        "probability_team_a_wins_by_3",
        "probability_team_a_wins_by_4",
        "probability_team_a_wins_by_5",
        "probability_team_a_wins_by_6_plus",
        "probability_team_b_wins_by_1",
        "probability_team_b_wins_by_2",
        "probability_team_b_wins_by_3",
        "probability_team_b_wins_by_4",
        "probability_team_b_wins_by_5",
        "probability_team_b_wins_by_6_plus",
        "probability_draw",
        "top_5_ev_decomposition",
        "top_10_ev_decomposition",
        "ev_explanation",
        "market_fair_btts_yes_probability",
        "market_fair_btts_no_probability",
        "model_implied_btts_yes_probability",
        "model_implied_btts_no_probability",
        "btts_fit_error",
        "btts_market_available",
        "model_probability_team_a_clean_sheet",
        "model_probability_team_b_clean_sheet",
        "model_probability_no_btts",
        "model_probability_btts",
        "extreme_favourite_audit_triggered",
        "top_clean_sheet_scores",
        "top_favourite_margin_probabilities",
        "top_5_favourite_score_count_probabilities",
        "high_score_cluster_scores",
        "three_four_five_nil_within_0_10_ev",
        "high_score_cluster",
        "high_score_cluster_close_alternatives",
        "current_final_recommendation",
        "normal_grid_poisson_recommendation",
        "larger_grid_poisson_recommendation",
        "normal_grid_recommendation",
        "larger_grid_recommendation",
        "recommendation_changes_with_larger_grid",
        "larger_grid_poisson_changes_recommendation",
        "larger_grid_differs_from_live_recommendation",
        "normal_grid_tail_mass",
        "larger_grid_tail_mass",
        "ev_favourite_3_0_normal_grid",
        "ev_favourite_4_0_normal_grid",
        "ev_favourite_5_0_normal_grid",
        "ev_favourite_3_0_larger_grid",
        "ev_favourite_4_0_larger_grid",
        "ev_favourite_5_0_larger_grid",
        "baseline_poisson_recommended_score",
        "baseline_poisson_ev_gap_best_vs_second",
        "correct_score_blended_recommended_score",
        "correct_score_blended_ev_gap_best_vs_second",
        "final_live_recommended_score",
        "final_live_ev_gap_best_vs_second",
        "model_recommendations_agree",
        "model_disagreement_warning",
        "dixon_coles_rho",
        "dixon_coles_rho_used",
        "dixon_coles_rho_source",
        "dixon_coles_rho_fit_error",
        "dixon_coles_recommended_score",
        "dixon_coles_best_expected_points",
        "dixon_coles_ev_gap_best_vs_second",
        "dixon_coles_top_5_ev_predictions",
        "dixon_coles_changes_recommendation",
        "market_consistent_recommended_score",
        "market_consistent_best_expected_points",
        "market_consistent_ev_gap_best_vs_second",
        "market_consistent_differs_from_default",
        "market_consistent_kl_divergence_vs_prior",
        "market_consistent_1x2_fit_error",
        "market_consistent_btts_fit_error",
        "market_consistent_total_goals_fit_error",
        "market_consistent_asian_handicap_fit_error",
        "market_consistent_correct_score_fit_error",
        "market_consistent_asian_totals_used",
        "market_consistent_asian_totals_shift_recommendation",
        "asian_handicap_lines_available",
        "asian_handicap_lines_used",
        "market_consistent_asian_handicap_lines_available",
        "market_consistent_asian_handicap_lines_selected",
        "market_consistent_asian_handicap_lines_skipped",
        "market_consistent_asian_handicap_lines_used",
        "market_consistent_asian_handicap_lines_skipped_detail",
        "market_consistent_asian_handicap_fit_error_selected",
        "market_consistent_asian_handicap_fit_error_all",
        "market_consistent_margin_distribution_before",
        "market_consistent_margin_distribution_after",
        "market_consistent_largest_margin_shift",
        "market_consistent_largest_margin_shift_value",
        "market_consistent_top_10_ev_scorelines",
        "market_consistent_top_10_ev_decomposition",
        "market_consistent_top_10_probability_scorelines",
        "market_consistent_poisson_ev_favourite_3_0",
        "market_consistent_poisson_ev_favourite_4_0",
        "market_consistent_poisson_ev_favourite_5_0",
        "market_consistent_matrix_ev_favourite_3_0",
        "market_consistent_matrix_ev_favourite_4_0",
        "market_consistent_matrix_ev_favourite_5_0",
        "estimated_most_crowded_public_score",
        "estimated_most_crowded_public_pick_share",
        "public_strategy_mode",
        "public_strategy_target",
        "public_field_size",
        "public_strategy_score",
        "public_strategy_ev_cost",
        "public_strategy_public_pick_share",
        "public_strategy_leverage_score",
        "public_strategy_reason",
        "friend_strategy_score",
        "friend_strategy_reason",
        "warnings",
        "number_of_bookmakers",
        "margin_removal_method",
        "has_over_under",
        "total_goals_lines_available",
        "total_goals_lines_used_for_calibration",
        "total_goals_lines_skipped_for_calibration",
        "total_goals_lines_skipped",
        "total_goals_line_fit_error",
        "total_goals_line_diagnostics",
        "over_under_2_5_used",
        "multi_line_totals_used",
        "has_btts",
        "has_qualification_odds",
        "has_correct_score_market",
        "correct_score_blend_applied",
        "correct_score_blend_suppressed",
        "correct_score_blend_note",
        "correct_score_scorelines_count",
        "correct_score_bookmakers_count",
        "correct_score_sparse_warning",
        "correct_score_aggregation_method",
        "number_of_correct_score_bookmakers",
        "average_correct_score_overround",
        "max_correct_score_overround",
        "total_goals_line_fit_error",
        "outlier_count",
        "top_outlier_examples",
        "scoreline_coverage_warning",
        "out_of_grid_scorelines_count",
        "correct_score_bookmaker_diagnostics",
        "selected_blend_weight",
        "effective_correct_score_poisson_weight",
        "kl_divergence",
        "top_10_market_scorelines",
        "top_10_blended_scorelines",
        "correct_score_poisson_weight",
        "correct_score_kl_divergence",
        "correct_score_market_top_10",
        "correct_score_blended_top_10",
        "ev_gap_best_vs_second",
        "ev_gap_best_vs_modal",
        "margin_ev_gap",
        "draw_vs_decisive_gap",
        "margin_diagnostic_note",
        "warning_flags",
        "lambda_a_near_bound",
        "lambda_b_near_bound",
    ]
    ordered_columns = list(dict.fromkeys(column for column in key_columns if column in frame))
    ordered_columns.extend(column for column in frame.columns if column not in ordered_columns)
    display_frame = frame[ordered_columns].rename(columns=aliases)
    ev_decomposition = _ev_decomposition_sheet(frame)
    high_score_diagnostics = _high_score_diagnostics_sheet(frame)
    market_consistent_diagnostics = _market_consistent_diagnostics_sheet(frame)
    with pd.ExcelWriter(path) as writer:
        display_frame.to_excel(writer, index=False, sheet_name="recommendations")
        if not ev_decomposition.empty:
            ev_decomposition.to_excel(writer, index=False, sheet_name="ev_decomposition")
        if not high_score_diagnostics.empty:
            high_score_diagnostics.to_excel(writer, index=False, sheet_name="high_score_diagnostics")
        if not market_consistent_diagnostics.empty:
            market_consistent_diagnostics.to_excel(writer, index=False, sheet_name="market_consistent")
        if asian_handicap_diagnostics is not None and not asian_handicap_diagnostics.empty:
            asian_handicap_diagnostics.to_excel(writer, index=False, sheet_name="asian_handicap")
        if margin_diagnostics is not None and not margin_diagnostics.empty:
            margin_diagnostics.to_excel(writer, index=False, sheet_name="margin_diagnostics")
        if margin_method_comparison is not None and not margin_method_comparison.empty:
            margin_method_comparison.to_excel(writer, index=False, sheet_name="margin_methods")
        if final_decision_dashboard is not None and not final_decision_dashboard.empty:
            final_decision_dashboard.to_excel(writer, index=False, sheet_name="final_decision_dashboard")

    probability_columns = {
        "market_a",
        "market_draw",
        "market_b",
        "favourite_probability",
        "model_a_win",
        "model_draw",
        "model_b_win",
        "market_fair_btts_yes_probability",
        "market_fair_btts_no_probability",
        "model_implied_btts_yes_probability",
        "model_implied_btts_no_probability",
        "model_probability_team_a_clean_sheet",
        "model_probability_team_b_clean_sheet",
        "model_probability_no_btts",
        "model_probability_btts",
        "probability_total_goals_0",
        "probability_total_goals_1",
        "probability_total_goals_2",
        "probability_total_goals_3",
        "probability_total_goals_4",
        "probability_total_goals_5",
        "probability_total_goals_6",
        "probability_total_goals_7_plus",
        "probability_total_goals_4_plus",
        "probability_team_a_scores_0",
        "probability_team_a_scores_1",
        "probability_team_a_scores_2",
        "probability_team_a_scores_3",
        "probability_team_a_scores_4",
        "probability_team_a_scores_5",
        "probability_team_a_scores_6_plus",
        "probability_team_b_scores_0",
        "probability_team_b_scores_1",
        "probability_team_b_scores_2",
        "probability_team_b_scores_3",
        "probability_team_b_scores_4",
        "probability_team_b_scores_5",
        "probability_team_b_scores_6_plus",
        "probability_team_a_wins_by_1",
        "probability_team_a_wins_by_2",
        "probability_team_a_wins_by_3",
        "probability_team_a_wins_by_4",
        "probability_team_a_wins_by_5",
        "probability_team_a_wins_by_6_plus",
        "probability_team_b_wins_by_1",
        "probability_team_b_wins_by_2",
        "probability_team_b_wins_by_3",
        "probability_team_b_wins_by_4",
        "probability_team_b_wins_by_5",
        "probability_team_b_wins_by_6_plus",
        "probability_draw",
        "exact_score_probability",
        "correct_goal_difference_probability",
        "correct_result_probability",
        "tail_probability_before_renormalisation",
        "estimated_most_crowded_public_pick_share",
        "public_strategy_public_pick_share",
        "public_strategy_exact_score_probability",
        "public_strategy_result_probability",
    }
    decimal_columns = {
        "lambda_a",
        "lambda_b",
        "calibration_loss",
        "correct_score_kl_divergence",
        "kl_divergence",
        "selected_blend_weight",
        "effective_correct_score_poisson_weight",
        "average_correct_score_overround",
        "max_correct_score_overround",
        "dixon_coles_rho",
        "dixon_coles_rho_used",
        "dixon_coles_rho_source",
        "dixon_coles_rho_fit_error",
        "public_strategy_leverage_score",
        "public_strategy_public_ranking_score",
        "expected_team_a_goals",
        "expected_team_b_goals",
        "expected_total_goals",
        "btts_fit_error",
        "normal_grid_tail_mass",
        "larger_grid_tail_mass",
        "ev_favourite_3_0_normal_grid",
        "ev_favourite_4_0_normal_grid",
        "ev_favourite_5_0_normal_grid",
        "ev_favourite_3_0_larger_grid",
        "ev_favourite_4_0_larger_grid",
        "ev_favourite_5_0_larger_grid",
        "market_consistent_kl_divergence_vs_prior",
        "market_consistent_1x2_fit_error",
        "market_consistent_btts_fit_error",
        "market_consistent_total_goals_fit_error",
        "market_consistent_asian_handicap_fit_error",
        "market_consistent_correct_score_fit_error",
        "margin_ev_gap",
        "draw_vs_decisive_gap",
        "best_draw_ev",
        "best_decisive_ev",
        "market_consistent_poisson_ev_favourite_3_0",
        "market_consistent_poisson_ev_favourite_4_0",
        "market_consistent_poisson_ev_favourite_5_0",
        "market_consistent_matrix_ev_favourite_3_0",
        "market_consistent_matrix_ev_favourite_4_0",
        "market_consistent_matrix_ev_favourite_5_0",
    }
    ev_columns = {
        "best_expected_points",
        "ev_gap_to_second",
        "ev_gap_to_third",
        "ev_gap_best_vs_second",
        "ev_gap_best_vs_modal",
        "baseline_poisson_ev_gap_best_vs_second",
        "correct_score_blended_ev_gap_best_vs_second",
        "final_live_ev_gap_best_vs_second",
        "dixon_coles_ev_gap_best_vs_second",
        "dixon_coles_best_expected_points",
        "market_consistent_best_expected_points",
        "market_consistent_ev_gap_best_vs_second",
        "public_strategy_ev_cost",
    }
    wrapped_columns = {
        "top_5_ev_predictions",
        "plausible_top_alternatives",
        "suppressed_ev_candidates",
        "suppression_reason",
        "top_5_ev_decomposition",
        "top_10_ev_decomposition",
        "ev_explanation",
        "close_alternatives",
        "plausible_alternatives",
        "decision_note",
        "risk_notes",
        "top_clean_sheet_scores",
        "top_favourite_margin_probabilities",
        "top_5_favourite_score_count_probabilities",
        "high_score_cluster_scores",
        "correct_score_market_top_10",
        "correct_score_blended_top_10",
        "top_10_market_scorelines",
        "top_10_blended_scorelines",
        "correct_score_sparse_warning",
        "correct_score_blend_note",
        "top_outlier_examples",
        "scoreline_coverage_warning",
        "correct_score_bookmaker_diagnostics",
        "warning_flags",
        "warnings",
        "total_goals_lines_available",
        "total_goals_lines_used_for_calibration",
        "total_goals_lines_skipped_for_calibration",
        "total_goals_lines_skipped",
        "total_goals_line_diagnostics",
        "model_disagreement_warning",
        "dixon_coles_top_5_ev_predictions",
        "market_consistent_asian_totals_used",
        "asian_handicap_lines_available",
        "asian_handicap_lines_used",
        "market_consistent_asian_handicap_lines_used",
        "market_consistent_margin_distribution_before",
        "market_consistent_margin_distribution_after",
        "margin_diagnostic_note",
        "market_consistent_top_10_ev_scorelines",
        "market_consistent_top_10_ev_decomposition",
        "market_consistent_top_10_probability_scorelines",
        "public_strategy_reason",
        "friend_strategy_reason",
    }
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    worksheet = workbook["recommendations"]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, cell in enumerate(worksheet[1], start=1):
        column_name = str(cell.value)
        for value_cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for value in value_cell:
                if column_name in probability_columns:
                    value.number_format = "0.00%"
                elif column_name in decimal_columns:
                    value.number_format = "0.0000"
                elif column_name in ev_columns:
                    value.number_format = "0.000"
                if column_name in wrapped_columns:
                    value.alignment = Alignment(wrap_text=True, vertical="top")
        contents = [str(cell.value or "")]
        contents.extend(str(row[0].value or "") for row in worksheet.iter_cols(min_col=index, max_col=index, min_row=2))
        max_length = max(len(content) for content in contents)
        maximum_width = 60 if column_name in wrapped_columns else 32
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 10), maximum_width)

    workbook.save(path)


def export_world_cup_submission_sheet_excel(frame: pd.DataFrame, path: str | Path) -> None:
    """Export a concise World Cup prediction-entry workbook."""

    path = Path(path)
    ensure_parent_directory(path)
    submission = frame.copy()
    submission["top_3_alternatives"] = submission["plausible_top_alternatives"].fillna("").map(
        lambda value: "; ".join(str(value).split("; ")[:3])
    )
    submission["notes"] = submission[["warning_flags", "warnings", "source_notes"]].fillna("").apply(
        lambda values: "; ".join(value for value in values if value),
        axis=1,
    )
    submission = submission[
        [
            "match_id",
            "date",
            "stage",
            "group",
            "team_a",
            "team_b",
            "recommended_score",
            "final_decision_score",
            "confidence_level",
            "manual_review_flag",
            "main_alternative_score",
            "plausible_alternatives",
            "decision_note",
            "recommended_qualifier",
            "best_expected_points",
            "recommendation_confidence",
            "plausible_top_alternatives",
            "suppressed_ev_candidates",
            "close_alternatives",
            "ev_gap_to_second",
            "ev_gap_to_third",
            "favourite_bucket",
            "top_3_alternatives",
            "notes",
        ]
    ].sort_values(["date", "match_id"], kind="stable")
    submission.to_excel(path, index=False, sheet_name="submission")

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    worksheet = workbook["submission"]
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    wrapped_columns = {
        "top_3_alternatives",
        "notes",
        "plausible_top_alternatives",
        "plausible_alternatives",
        "suppressed_ev_candidates",
        "main_alternative_score",
        "close_alternatives",
        "decision_note",
    }
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, cell in enumerate(worksheet[1], start=1):
        column_name = str(cell.value)
        for value_cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for value in value_cell:
                if column_name in {"best_expected_points", "ev_gap_to_second", "ev_gap_to_third"}:
                    value.number_format = "0.000"
                if column_name in wrapped_columns:
                    value.alignment = Alignment(wrap_text=True, vertical="top")
        contents = [str(cell.value or "")]
        contents.extend(str(row[0].value or "") for row in worksheet.iter_cols(min_col=index, max_col=index, min_row=2))
        max_length = max(len(content) for content in contents)
        maximum_width = 60 if column_name in wrapped_columns else 32
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 10), maximum_width)

    workbook.save(path)


def export_report_bundle(
    match_report: pd.DataFrame,
    friend_report: pd.DataFrame,
    output_dir: str | Path,
) -> None:
    """Write the standard pre-match reports to CSV and one Excel workbook."""

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    match_report.to_csv(output_dir / "match_recommendations.csv", index=False)
    friend_report.to_csv(output_dir / "friend_prediction_ev.csv", index=False)
    with pd.ExcelWriter(output_dir / "prediction_reports.xlsx") as writer:
        match_report.to_excel(writer, sheet_name="recommendations", index=False)
        friend_report.to_excel(writer, sheet_name="friend_ev", index=False)


def export_standings_bundle(
    scored_predictions: pd.DataFrame,
    standings: pd.DataFrame,
    output_dir: str | Path,
) -> None:
    """Write realised match scoring and standings reports."""

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    scored_predictions.to_csv(output_dir / "realised_prediction_scores.csv", index=False)
    standings.to_csv(output_dir / "standings.csv", index=False)
    with pd.ExcelWriter(output_dir / "standings_reports.xlsx") as writer:
        scored_predictions.to_excel(writer, sheet_name="prediction_scores", index=False)
        standings.to_excel(writer, sheet_name="standings", index=False)
