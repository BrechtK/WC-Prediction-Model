"""Compare live recommendation sensitivity across correct-score blend weights."""

from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path
from collections.abc import Sequence

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wc_predictor.config import ProjectConfig
from wc_predictor.market_data import load_correct_score_odds, load_odds, load_total_goals_odds
from wc_predictor.paths import (
    CACHE_CORE_ODDS_PATH,
    CACHE_CORRECT_SCORE_ODDS_PATH,
    OUTPUT_WEIGHT_SENSITIVITY_XLSX_PATH,
)
from wc_predictor.utils import ensure_parent_directory
from wc_predictor.workflow import run_prediction_workflow

DEFAULT_WORLD_CUP_ODDS_PATH = CACHE_CORE_ODDS_PATH
DEFAULT_CORRECT_SCORE_ODDS_PATH = CACHE_CORRECT_SCORE_ODDS_PATH
DEFAULT_OUTPUT_PATH = OUTPUT_WEIGHT_SENSITIVITY_XLSX_PATH
DEFAULT_CORRECT_SCORE_WEIGHTS = tuple(round(value / 100, 2) for value in range(100, -1, -5))
SUMMARY_WEIGHTS = (1.00, 0.85, 0.75, 0.50, 0.00)

DETAIL_COLUMNS = [
    "match_id",
    "team_a",
    "team_b",
    "weight",
    "recommended_score",
    "recommended_qualifier",
    "best_expected_points",
    "most_likely_scoreline",
    "top_5_ev_predictions",
    "top_10_market_scorelines",
    "top_10_blended_scorelines",
    "correct_score_aggregation_method",
    "correct_score_bookmakers_count",
    "average_correct_score_overround",
    "max_correct_score_overround",
    "correct_score_outlier_count",
    "kl_market_vs_poisson",
    "warning_flags",
]

@dataclass(frozen=True)
class CorrectScoreWeightComparison:
    """Detailed weight runs and match-level recommendation sensitivity."""

    details: pd.DataFrame
    sensitivity: pd.DataFrame


def parse_weight_grid(value: str) -> tuple[float, ...]:
    """Parse a comma-separated CLI weight grid."""

    try:
        return validate_weight_grid(tuple(float(item.strip()) for item in value.split(",") if item.strip()))
    except ValueError as error:
        raise ValueError("--weights must be comma-separated numbers between 0 and 1") from error


def validate_weight_grid(weights: Sequence[float]) -> tuple[float, ...]:
    """Validate, deduplicate, and return weights from pure Poisson downward."""

    numeric = tuple(float(weight) for weight in weights)
    if not numeric:
        raise ValueError("At least one correct-score blend weight is required")
    if not np.isfinite(numeric).all() or any(weight < 0 or weight > 1 for weight in numeric):
        raise ValueError("Correct-score blend weights must lie between zero and one")
    return tuple(sorted(set(numeric), reverse=True))


def _weight_column(weight: float) -> str:
    return f"recommendation_at_w_{weight:.2f}".replace(".", "_")


SENSITIVITY_COLUMNS = [
    "match_id",
    "team_a",
    "team_b",
    *[_weight_column(weight) for weight in SUMMARY_WEIGHTS],
    "unique_recommended_scores_across_weights",
    "recommendation_changes",
    "first_weight_where_recommendation_changes_from_poisson",
    "stable_weight_range_around_0_85",
    "ev_gap_at_w_0_85",
    "notes",
]


def _row_at_weight(group: pd.DataFrame, weight: float) -> pd.Series | None:
    rows = group[np.isclose(group["weight"], weight)]
    return rows.iloc[0] if not rows.empty else None


def _score_at_weight(group: pd.DataFrame, weight: float) -> object:
    row = _row_at_weight(group, weight)
    return row["recommended_score"] if row is not None else pd.NA


def _first_change_from_poisson(group: pd.DataFrame) -> object:
    poisson = _row_at_weight(group, 1.0)
    if poisson is None:
        return pd.NA
    for _, row in group.sort_values("weight", ascending=False).iterrows():
        if row["weight"] < 1.0 and row["recommended_score"] != poisson["recommended_score"]:
            return float(row["weight"])
    return pd.NA


def _stable_weight_range_around(group: pd.DataFrame, anchor_weight: float = 0.85) -> object:
    ordered = group.sort_values("weight", ascending=False).reset_index(drop=True)
    anchor_rows = ordered.index[np.isclose(ordered["weight"], anchor_weight)]
    if len(anchor_rows) == 0:
        return pd.NA
    anchor_index = int(anchor_rows[0])
    anchor_score = ordered.loc[anchor_index, "recommended_score"]
    lower_index = anchor_index
    upper_index = anchor_index
    while lower_index > 0 and ordered.loc[lower_index - 1, "recommended_score"] == anchor_score:
        lower_index -= 1
    while upper_index + 1 < len(ordered) and ordered.loc[upper_index + 1, "recommended_score"] == anchor_score:
        upper_index += 1
    values = ordered.loc[[lower_index, upper_index], "weight"].astype(float)
    return f"[{values.min():.2f}, {values.max():.2f}]"


def _summary_notes(group: pd.DataFrame) -> str:
    notes: list[str] = []
    if group["recommended_score"].nunique() == 1:
        notes.append("stable_across_all_weights")
    qualifiers = {str(value) for value in group["recommended_qualifier"].fillna("") if str(value)}
    if len(qualifiers) > 1:
        notes.append("qualifier_changes_across_weights")
    if _row_at_weight(group, 1.0) is None:
        notes.append("pure_poisson_weight_not_in_grid")
    if _row_at_weight(group, 0.85) is None:
        notes.append("w_0_85_not_in_grid")
    return "; ".join(notes)


def _build_sensitivity_summary(details: pd.DataFrame) -> pd.DataFrame:
    """Return one row per match with reference weights and transition diagnostics."""

    rows: list[dict[str, object]] = []
    for match_id, group in details.groupby("match_id", sort=True):
        first = group.iloc[0]
        unique_scores = sorted(set(group["recommended_score"].astype(str)))
        row: dict[str, object] = {
            "match_id": match_id,
            "team_a": first["team_a"],
            "team_b": first["team_b"],
        }
        for weight in SUMMARY_WEIGHTS:
            row[_weight_column(weight)] = _score_at_weight(group, weight)
        row.update(
            {
                "unique_recommended_scores_across_weights": "; ".join(unique_scores),
                "recommendation_changes": "yes" if len(unique_scores) > 1 else "no",
                "first_weight_where_recommendation_changes_from_poisson": _first_change_from_poisson(group),
                "stable_weight_range_around_0_85": _stable_weight_range_around(group),
                "ev_gap_at_w_0_85": (
                    _row_at_weight(group, 0.85)["ev_gap_best_vs_second"]
                    if _row_at_weight(group, 0.85) is not None
                    else pd.NA
                ),
                "notes": _summary_notes(group),
            }
        )
        rows.append(row)
    summary = pd.DataFrame(rows, columns=SENSITIVITY_COLUMNS)
    return summary.sort_values(["recommendation_changes", "match_id"], ascending=[False, True]).reset_index(drop=True)


def compare_correct_score_weights(
    odds_path: str | Path = DEFAULT_WORLD_CUP_ODDS_PATH,
    correct_score_odds_path: str | Path = DEFAULT_CORRECT_SCORE_ODDS_PATH,
    *,
    weights: Sequence[float] = DEFAULT_CORRECT_SCORE_WEIGHTS,
    config: ProjectConfig | None = None,
    total_goals_odds_path: str | Path | None = None,
) -> CorrectScoreWeightComparison:
    """Run the existing recommendation workflow across one live weight grid."""

    odds = load_odds(odds_path)
    correct_score_odds = load_correct_score_odds(correct_score_odds_path)
    total_goals_odds = load_total_goals_odds(total_goals_odds_path) if total_goals_odds_path is not None else None
    base_config = config or ProjectConfig()
    rows: list[pd.DataFrame] = []
    for weight in validate_weight_grid(weights):
        weighted_config = replace(base_config, correct_score_poisson_weight=weight)
        report = run_prediction_workflow(
            odds,
            config=weighted_config,
            correct_score_odds=correct_score_odds,
            total_goals_odds=total_goals_odds,
        ).match_report
        report = report[report["has_correct_score_market"]].copy()
        report["weight"] = weight
        report["correct_score_outlier_count"] = report["outlier_count"]
        report["kl_market_vs_poisson"] = report["correct_score_kl_divergence"]
        rows.append(report[[*DETAIL_COLUMNS, "ev_gap_best_vs_second"]])
    details = pd.concat(rows, ignore_index=True)
    details = details.sort_values(["match_id", "weight"], ascending=[True, False], kind="stable").reset_index(drop=True)
    return CorrectScoreWeightComparison(details, _build_sensitivity_summary(details))


def _style_worksheet(
    worksheet,
    *,
    decimal_columns: set[str] | None = None,
    wrapped_columns: set[str] | None = None,
) -> None:
    """Apply compact formatting to one workbook sheet."""

    decimal_columns = decimal_columns or set()
    wrapped_columns = wrapped_columns or set()
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, cell in enumerate(worksheet[1], start=1):
        column_name = str(cell.value)
        for values in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for value in values:
                if column_name in decimal_columns:
                    value.number_format = "0.0000"
                if column_name in wrapped_columns:
                    value.alignment = Alignment(wrap_text=True, vertical="top")
        contents = [str(cell.value or "")]
        for values in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
            contents.extend(str(value.value or "") for value in values)
        maximum_width = 60 if column_name in wrapped_columns else 38
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(max(len(content) for content in contents) + 2, 10),
            maximum_width,
        )


def export_correct_score_weight_comparison(
    comparison: CorrectScoreWeightComparison,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> Path:
    """Write detailed weight runs and match-level sensitivity to Excel."""

    output_path = Path(output_path)
    ensure_parent_directory(output_path)
    details = comparison.details[DETAIL_COLUMNS]
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        details.to_excel(writer, index=False, sheet_name="details")
        comparison.sensitivity.to_excel(writer, index=False, sheet_name="sensitivity")

    from openpyxl import load_workbook

    workbook = load_workbook(output_path)
    _style_worksheet(
        workbook["details"],
        decimal_columns={
            "weight",
            "best_expected_points",
            "average_correct_score_overround",
            "max_correct_score_overround",
            "kl_market_vs_poisson",
        },
        wrapped_columns={
            "top_5_ev_predictions",
            "top_10_market_scorelines",
            "top_10_blended_scorelines",
            "warning_flags",
        },
    )
    _style_worksheet(
        workbook["sensitivity"],
        decimal_columns={
            "first_weight_where_recommendation_changes_from_poisson",
            "ev_gap_at_w_0_85",
        },
        wrapped_columns={
            "unique_recommended_scores_across_weights",
            "notes",
        },
    )
    workbook.save(output_path)
    return output_path


def format_correct_score_weight_comparison_summary(
    comparison: CorrectScoreWeightComparison,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> str:
    """Render a concise live blend-weight sensitivity summary."""

    summary = comparison.sensitivity
    changed = summary["recommendation_changes"] == "yes"
    stable = ~changed
    poisson = summary[_weight_column(1.00)]
    at_085 = summary[_weight_column(0.85)]
    at_050 = summary[_weight_column(0.50)]
    differs_at_085 = poisson.notna() & at_085.notna() & (at_085 != poisson)
    differs_at_050 = poisson.notna() & at_050.notna() & (at_050 != poisson)
    return "\n".join(
        [
            "Correct-score blend-weight sensitivity",
            "--------------------------------------",
            f"Matches compared: {len(summary)}",
            f"Matches where recommendation changes across weights: {int(changed.sum())}",
            f"Matches stable for all weights: {int(stable.sum())}",
            f"Matches where w=0.85 differs from pure Poisson: {int(differs_at_085.sum())}",
            f"Matches where w=0.50 differs from pure Poisson: {int(differs_at_050.sum())}",
            f"Output workbook path: {Path(output_path)}",
        ]
    )
