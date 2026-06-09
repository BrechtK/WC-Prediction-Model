"""Compare Dixon-Coles challenger sensitivity across a configurable rho grid."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, replace
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wc_predictor.config import ProjectConfig
from wc_predictor.live_prediction import LivePredictionSettings, parse_live_prediction_inputs
from wc_predictor.market_data import load_correct_score_odds, load_odds, load_total_goals_odds
from wc_predictor.paths import OUTPUT_DIXON_COLES_SENSITIVITY_XLSX_PATH
from wc_predictor.utils import ensure_parent_directory
from wc_predictor.workflow import run_prediction_workflow

DEFAULT_DIXON_COLES_RHOS = (-0.20, -0.15, -0.10, -0.05, 0.00, 0.05, 0.10, 0.15, 0.20)
DEFAULT_OUTPUT_PATH = OUTPUT_DIXON_COLES_SENSITIVITY_XLSX_PATH

DETAIL_COLUMNS = [
    "match_id",
    "team_a",
    "team_b",
    "rho",
    "baseline_poisson_recommended_score",
    "dixon_coles_recommended_score",
    "dixon_coles_best_expected_points",
    "dixon_coles_top_5_ev_predictions",
    "final_live_recommended_score",
    "dixon_coles_changes_recommendation",
    "p_dc_0_0",
    "p_dc_1_0",
    "p_dc_0_1",
    "p_dc_1_1",
    "dixon_coles_matrix_probability_sum",
    "dixon_coles_matrix_normalised",
    "warning_flags",
]

SUMMARY_COLUMNS = [
    "match_id",
    "team_a",
    "team_b",
    "recommendation_at_rho_0",
    "unique_dixon_coles_recommendations_across_rho",
    "recommendation_changes_across_rho",
    "first_rho_where_recommendation_changes_from_rho_0",
    "stable_rho_interval_around_0",
    "final_live_recommended_score",
    "live_recommendation_agreement_count",
    "dixon_coles_runs",
    "live_recommendation_agrees_with_all_dixon_coles",
    "live_recommendation_agrees_with_most_dixon_coles",
    "notes",
]


@dataclass(frozen=True)
class DixonColesRhoComparison:
    """Detailed rho runs and match-level Dixon-Coles sensitivity."""

    details: pd.DataFrame
    summary: pd.DataFrame


def parse_rho_grid(value: str) -> tuple[float, ...]:
    """Parse a comma-separated CLI rho grid."""

    try:
        return validate_rho_grid(tuple(float(item.strip()) for item in value.split(",") if item.strip()))
    except ValueError as error:
        raise ValueError("--rhos must be comma-separated finite numbers") from error


def validate_rho_grid(rhos: Sequence[float]) -> tuple[float, ...]:
    """Validate, deduplicate, and return rho values in ascending order."""

    numeric = tuple(float(rho) for rho in rhos)
    if not numeric:
        raise ValueError("At least one Dixon-Coles rho value is required")
    if not np.isfinite(numeric).all():
        raise ValueError("Dixon-Coles rho values must be finite")
    return tuple(sorted(set(numeric)))


def _row_at_rho(group: pd.DataFrame, rho: float) -> pd.Series | None:
    rows = group[np.isclose(group["rho"], rho)]
    return rows.iloc[0] if not rows.empty else None


def _first_change_from_zero(group: pd.DataFrame) -> object:
    zero = _row_at_rho(group, 0.0)
    if zero is None:
        return pd.NA
    changed = group[group["dixon_coles_recommended_score"] != zero["dixon_coles_recommended_score"]].copy()
    if changed.empty:
        return pd.NA
    changed["distance_from_zero"] = changed["rho"].abs()
    return float(changed.sort_values(["distance_from_zero", "rho"], kind="stable").iloc[0]["rho"])


def _stable_rho_interval_around_zero(group: pd.DataFrame) -> object:
    ordered = group.sort_values("rho", kind="stable").reset_index(drop=True)
    zero_rows = ordered.index[np.isclose(ordered["rho"], 0.0)]
    if len(zero_rows) == 0:
        return pd.NA
    zero_index = int(zero_rows[0])
    zero_score = ordered.loc[zero_index, "dixon_coles_recommended_score"]
    lower_index = zero_index
    upper_index = zero_index
    while lower_index > 0 and ordered.loc[lower_index - 1, "dixon_coles_recommended_score"] == zero_score:
        lower_index -= 1
    while (
        upper_index + 1 < len(ordered)
        and ordered.loc[upper_index + 1, "dixon_coles_recommended_score"] == zero_score
    ):
        upper_index += 1
    return f"[{ordered.loc[lower_index, 'rho']:.2f}, {ordered.loc[upper_index, 'rho']:.2f}]"


def _summary_notes(group: pd.DataFrame) -> str:
    notes: list[str] = []
    if group["dixon_coles_recommended_score"].nunique() == 1:
        notes.append("stable_across_all_rho_values")
    if _row_at_rho(group, 0.0) is None:
        notes.append("rho_0_not_in_grid")
    if group["final_live_recommended_score"].nunique() > 1:
        notes.append("unexpected_final_live_recommendation_change")
    if not group["dixon_coles_matrix_normalised"].all():
        notes.append("non_normalised_dixon_coles_matrix")
    return "; ".join(notes)


def _build_rho_summary(details: pd.DataFrame) -> pd.DataFrame:
    """Return one sensitivity row per match."""

    rows: list[dict[str, object]] = []
    for match_id, group in details.groupby("match_id", sort=True):
        first = group.iloc[0]
        zero = _row_at_rho(group, 0.0)
        unique_scores = sorted(set(group["dixon_coles_recommended_score"].astype(str)))
        live_score = str(first["final_live_recommended_score"])
        agreement_count = int((group["dixon_coles_recommended_score"].astype(str) == live_score).sum())
        rows.append(
            {
                "match_id": match_id,
                "team_a": first["team_a"],
                "team_b": first["team_b"],
                "recommendation_at_rho_0": zero["dixon_coles_recommended_score"] if zero is not None else pd.NA,
                "unique_dixon_coles_recommendations_across_rho": "; ".join(unique_scores),
                "recommendation_changes_across_rho": "yes" if len(unique_scores) > 1 else "no",
                "first_rho_where_recommendation_changes_from_rho_0": _first_change_from_zero(group),
                "stable_rho_interval_around_0": _stable_rho_interval_around_zero(group),
                "final_live_recommended_score": live_score,
                "live_recommendation_agreement_count": agreement_count,
                "dixon_coles_runs": len(group),
                "live_recommendation_agrees_with_all_dixon_coles": agreement_count == len(group),
                "live_recommendation_agrees_with_most_dixon_coles": agreement_count > len(group) / 2,
                "notes": _summary_notes(group),
            }
        )
    summary = pd.DataFrame(rows, columns=SUMMARY_COLUMNS)
    return summary.sort_values(["recommendation_changes_across_rho", "match_id"], ascending=[False, True]).reset_index(
        drop=True
    )


def compare_dixon_coles_rhos(
    odds_path: str | Path,
    *,
    correct_score_odds_path: str | Path | None = None,
    total_goals_odds_path: str | Path | None = None,
    rhos: Sequence[float] = DEFAULT_DIXON_COLES_RHOS,
    config: ProjectConfig | None = None,
) -> DixonColesRhoComparison:
    """Run the existing recommendation workflow across a Dixon-Coles rho grid."""

    odds = load_odds(odds_path)
    correct_score_odds = load_correct_score_odds(correct_score_odds_path) if correct_score_odds_path is not None else None
    total_goals_odds = load_total_goals_odds(total_goals_odds_path) if total_goals_odds_path is not None else None
    base_config = config or ProjectConfig(correct_score_poisson_weight=0.85)
    rows: list[dict[str, object]] = []
    for rho in validate_rho_grid(rhos):
        rho_config = replace(base_config, dixon_coles_rho=rho, enable_dixon_coles_rho_estimation=False)
        workflow = run_prediction_workflow(
            odds,
            config=rho_config,
            correct_score_odds=correct_score_odds,
            total_goals_odds=total_goals_odds,
        )
        report = workflow.match_report.set_index("match_id")
        for match_id, matrix_by_model in workflow.challenger_score_matrices.items():
            row = report.loc[match_id]
            matrix = matrix_by_model["dixon_coles"]
            rows.append(
                {
                    "match_id": match_id,
                    "team_a": row["team_a"],
                    "team_b": row["team_b"],
                    "rho": rho,
                    "baseline_poisson_recommended_score": row["baseline_poisson_recommended_score"],
                    "dixon_coles_recommended_score": row["dixon_coles_recommended_score"],
                    "dixon_coles_best_expected_points": row["dixon_coles_best_expected_points"],
                    "dixon_coles_top_5_ev_predictions": row["dixon_coles_top_5_ev_predictions"],
                    "final_live_recommended_score": row["final_live_recommended_score"],
                    "dixon_coles_changes_recommendation": row["dixon_coles_changes_recommendation"],
                    "p_dc_0_0": matrix.exact_score_probability(0, 0),
                    "p_dc_1_0": matrix.exact_score_probability(1, 0),
                    "p_dc_0_1": matrix.exact_score_probability(0, 1),
                    "p_dc_1_1": matrix.exact_score_probability(1, 1),
                    "dixon_coles_matrix_probability_sum": matrix.grid_probability,
                    "dixon_coles_matrix_normalised": bool(np.isclose(matrix.grid_probability, 1.0, atol=1e-9)),
                    "warning_flags": row["warning_flags"],
                }
            )
    details = pd.DataFrame(rows, columns=DETAIL_COLUMNS)
    details = details.sort_values(["match_id", "rho"], kind="stable").reset_index(drop=True)
    return DixonColesRhoComparison(details, _build_rho_summary(details))


def compare_live_dixon_coles_rhos(
    settings: LivePredictionSettings | None = None,
    *,
    rhos: Sequence[float] = DEFAULT_DIXON_COLES_RHOS,
    config: ProjectConfig | None = None,
) -> DixonColesRhoComparison:
    """Parse the current live OddsPortal pastes and run rho sensitivity."""

    parsed = parse_live_prediction_inputs(settings)
    return compare_dixon_coles_rhos(
        parsed.settings.core_odds_output_path,
        correct_score_odds_path=parsed.settings.correct_score_output_path if parsed.has_correct_scores else None,
        total_goals_odds_path=parsed.settings.total_goals_output_path if parsed.has_total_goals else None,
        rhos=rhos,
        config=config,
    )


def _style_worksheet(
    worksheet,
    *,
    decimal_columns: set[str] | None = None,
    wrapped_columns: set[str] | None = None,
) -> None:
    """Apply compact formatting to one workbook sheet."""

    decimal_columns = decimal_columns or set()
    wrapped_columns = wrapped_columns or set()
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, cell in enumerate(worksheet[1], start=1):
        column_name = str(cell.value)
        for values in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
            for value in values:
                if column_name in decimal_columns:
                    value.number_format = "0.0000"
                if column_name in wrapped_columns:
                    value.alignment = Alignment(wrap_text=True, vertical="top")
        contents = [str(cell.value or "")]
        for values in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
            contents.extend(str(value.value or "") for value in values)
        maximum_width = 60 if column_name in wrapped_columns else 42
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(max(len(content) for content in contents) + 2, 10),
            maximum_width,
        )


def export_dixon_coles_rho_comparison(
    comparison: DixonColesRhoComparison,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> Path:
    """Write detailed rho runs and match-level sensitivity to Excel."""

    output_path = Path(output_path)
    ensure_parent_directory(output_path)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        comparison.details.to_excel(writer, index=False, sheet_name="details")
        comparison.summary.to_excel(writer, index=False, sheet_name="summary")

    from openpyxl import load_workbook

    workbook = load_workbook(output_path)
    _style_worksheet(
        workbook["details"],
        decimal_columns={
            "rho",
            "dixon_coles_best_expected_points",
            "p_dc_0_0",
            "p_dc_1_0",
            "p_dc_0_1",
            "p_dc_1_1",
            "dixon_coles_matrix_probability_sum",
        },
        wrapped_columns={"dixon_coles_top_5_ev_predictions", "warning_flags"},
    )
    _style_worksheet(
        workbook["summary"],
        decimal_columns={"first_rho_where_recommendation_changes_from_rho_0"},
        wrapped_columns={"unique_dixon_coles_recommendations_across_rho", "notes"},
    )
    workbook.save(output_path)
    return output_path


def format_dixon_coles_rho_comparison_summary(
    comparison: DixonColesRhoComparison,
    output_path: str | Path = DEFAULT_OUTPUT_PATH,
) -> str:
    """Render a concise Dixon-Coles rho sensitivity summary."""

    summary = comparison.summary
    changed = summary["recommendation_changes_across_rho"] == "yes"
    differs_from_live = ~summary["live_recommendation_agrees_with_all_dixon_coles"]
    return "\n".join(
        [
            "Dixon-Coles challenger rho sensitivity",
            "--------------------------------------",
            f"Matches compared: {len(summary)}",
            f"Matches where Dixon-Coles recommendation changes across rho: {int(changed.sum())}",
            f"Matches where Dixon-Coles differs from final live recommendation: {int(differs_from_live.sum())}",
            f"Output workbook path: {Path(output_path)}",
        ]
    )
