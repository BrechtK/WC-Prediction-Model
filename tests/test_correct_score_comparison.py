from pathlib import Path
import runpy
import sys

from openpyxl import load_workbook

from wc_predictor.correct_score_comparison import (
    CORRECT_SCORE_AGGREGATION_METHODS,
    DETAIL_COLUMNS,
    compare_correct_score_aggregation_methods,
    export_correct_score_aggregation_comparison,
    format_correct_score_aggregation_comparison_summary,
)


EXAMPLES = Path("data/examples")
RUN_SCRIPT = Path("scripts/compare_correct_score_aggregation_methods.py").resolve()


def _comparison():
    return compare_correct_score_aggregation_methods(
        EXAMPLES / "example_world_cup_odds.csv",
        EXAMPLES / "example_world_cup_correct_score_odds.csv",
    )


def test_correct_score_aggregation_comparison_runs_all_requested_methods() -> None:
    comparison = _comparison()

    assert comparison.details.columns.tolist() == DETAIL_COLUMNS
    assert set(comparison.details["aggregation_method"]) == set(CORRECT_SCORE_AGGREGATION_METHODS)
    assert len(comparison.details) == 12
    assert {
        "match_id",
        "team_a",
        "team_b",
        *CORRECT_SCORE_AGGREGATION_METHODS,
        "unique_recommended_scores",
        "number_of_unique_recommended_scores",
        "recommended_score_changes",
        "aggregation_method_sensitive",
    }.issubset(comparison.sensitivity.columns)
    assert len(comparison.sensitivity) == 2
    assert comparison.sensitivity["number_of_unique_recommended_scores"].ge(1).all()


def test_correct_score_aggregation_comparison_exports_formatted_workbook(tmp_path: Path) -> None:
    output_path = tmp_path / "comparison.xlsx"

    export_correct_score_aggregation_comparison(_comparison(), output_path)

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["details", "sensitivity"]
    details = workbook["details"]
    sensitivity = workbook["sensitivity"]
    detail_headers = {cell.value: index + 1 for index, cell in enumerate(details[1])}
    assert details.freeze_panes == "A2"
    assert sensitivity.freeze_panes == "A2"
    assert details.cell(2, detail_headers["best_expected_points"]).number_format == "0.0000"
    assert details.cell(2, detail_headers["top_5_ev_predictions"]).alignment.wrap_text


def test_correct_score_aggregation_comparison_console_summary_is_compact() -> None:
    summary = format_correct_score_aggregation_comparison_summary(_comparison())

    assert "Correct-Score Aggregation Comparison" in summary
    assert "- Matches compared: 2" in summary
    assert "- Matches sensitive to aggregation method:" in summary
    assert "WC001" in summary
    assert "Alpha vs Beta" in summary


def test_correct_score_aggregation_comparison_script_runs_with_overrides(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    output_path = tmp_path / "comparison.xlsx"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "compare_correct_score_aggregation_methods.py",
            "--odds",
            str(EXAMPLES.resolve() / "example_world_cup_odds.csv"),
            "--correct-score-odds",
            str(EXAMPLES.resolve() / "example_world_cup_correct_score_odds.csv"),
            "--output",
            str(output_path),
        ],
    )

    runpy.run_path(str(RUN_SCRIPT), run_name="__main__")

    assert output_path.exists()
    assert "Correct-Score Aggregation Comparison" in capsys.readouterr().out
