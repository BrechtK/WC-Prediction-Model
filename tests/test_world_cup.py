import json
from pathlib import Path
import runpy
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

import wc_predictor.workflow as workflow_module
from wc_predictor.config import ProjectConfig
from wc_predictor.reporting import (
    format_world_cup_console_summary,
    format_world_cup_model_risk_summary,
    format_world_cup_prediction_diagnostics,
)
from wc_predictor.world_cup import (
    WORLD_CUP_ODDS_MISSING_MESSAGE,
    WorldCupPredictionSettings,
    create_world_cup_odds_file,
    resolve_world_cup_odds_input,
    run_world_cup_predictions,
)
from wc_predictor.workflow import (
    _btts_conflict_diagnostic,
    _btts_warning_flags,
    _build_final_decision_dashboard,
    _decision_aid,
    _blowout_risk_diagnostic,
    _draw_prone_diagnostic,
    _ev_explanation,
    _modal_draw_challenger_diagnostic,
    _plausible_alternative_decision_layer,
    _should_run_market_consistent_challenger,
    run_prediction_workflow,
)
from wc_predictor.optimiser import GroupPredictionEvaluation, GroupPredictionRecommendation


EXAMPLES = Path("data/examples")
TEMPLATES = Path("data/templates")
RUN_SCRIPT = Path("scripts/run_world_cup_predictions.py").resolve()
CREATE_SCRIPT = Path("scripts/create_world_cup_odds_file.py").resolve()
ODDS_COLUMNS = [
    "match_id",
    "date",
    "stage",
    "group",
    "team_a",
    "team_b",
    "bookmaker",
    "odds_a_win",
    "odds_draw",
    "odds_b_win",
    "odds_over_2_5",
    "odds_under_2_5",
    "odds_btts_yes",
    "odds_btts_no",
    "odds_a_qualifies",
    "odds_b_qualifies",
    "odds_timestamp",
    "odds_source_url",
    "source_quality",
    "notes",
]


def _settings(tmp_path: Path) -> WorldCupPredictionSettings:
    return WorldCupPredictionSettings(
        input_path=EXAMPLES / "example_world_cup_odds.csv",
        csv_output_path=tmp_path / "world_cup_recommendations.csv",
        xlsx_output_path=tmp_path / "world_cup_recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "world_cup_submission_sheet.xlsx",
    )


def test_world_cup_templates_have_expected_columns() -> None:
    csv_odds = pd.read_csv(TEMPLATES / "world_cup_odds_template.csv")
    xlsx_odds = pd.read_excel(TEMPLATES / "world_cup_odds_template.xlsx")
    predictions = pd.read_csv(TEMPLATES / "friend_predictions_template.csv")

    assert csv_odds.columns.tolist() == ODDS_COLUMNS
    assert xlsx_odds.columns.tolist() == ODDS_COLUMNS
    assert predictions.columns.tolist() == [
        "match_id",
        "player",
        "predicted_team_a_goals",
        "predicted_team_b_goals",
        "predicted_qualifier",
    ]


def test_world_cup_workflow_exports_real_tournament_recommendations(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    workflow = run_world_cup_predictions(settings)

    assert settings.csv_output_path.exists()
    assert settings.xlsx_output_path.exists()
    assert settings.submission_xlsx_output_path.exists()
    assert len(workflow.match_report) == 2
    assert workflow.match_report.loc[0, "group"] == "Group A"
    assert workflow.match_report.loc[1, "recommended_qualifier"] in {"Gamma", "Delta"}
    assert {
        "market_a_win",
        "market_draw",
        "market_b_win",
        "lambda_a",
        "lambda_b",
        "favourite_bucket",
        "most_likely_scoreline",
        "recommended_score",
        "top_5_ev_predictions",
        "plausible_top_alternatives",
        "suppressed_ev_candidates",
        "suppression_reason",
        "odds_timestamp_min",
        "odds_timestamp_max",
        "bookmakers_used",
        "number_of_bookmakers",
        "margin_removal_method",
        "has_over_under",
        "has_btts",
        "has_qualification_odds",
        "favourite_probability",
        "ev_gap_best_vs_second",
        "ev_gap_best_vs_modal",
        "baseline_poisson_recommended_score",
        "baseline_poisson_ev_gap_best_vs_second",
        "correct_score_blended_recommended_score",
        "correct_score_blended_ev_gap_best_vs_second",
        "final_live_recommended_score",
        "final_live_ev_gap_best_vs_second",
        "model_recommendations_agree",
        "model_disagreement_warning",
        "dixon_coles_rho",
        "dixon_coles_rho_used",
        "dixon_coles_rho_source",
        "dixon_coles_rho_fit_error",
        "dixon_coles_recommended_score",
        "dixon_coles_ev_gap_best_vs_second",
        "dixon_coles_top_5_ev_predictions",
        "dixon_coles_changes_recommendation",
        "market_consistent_recommended_score",
        "market_consistent_best_expected_points",
        "market_consistent_ev_gap_best_vs_second",
        "market_consistent_top_10_ev_scorelines",
        "market_consistent_top_10_probability_scorelines",
        "market_consistent_kl_divergence_vs_prior",
        "market_consistent_1x2_fit_error",
        "market_consistent_btts_fit_error",
        "market_consistent_total_goals_fit_error",
        "market_consistent_correct_score_fit_error",
        "market_consistent_asian_totals_used",
        "market_consistent_differs_from_default",
        "market_consistent_asian_totals_shift_recommendation",
        "estimated_most_crowded_public_score",
        "estimated_most_crowded_public_pick_share",
        "public_strategy_score",
        "public_strategy_ev_cost",
        "public_strategy_public_pick_share",
        "public_strategy_leverage_score",
        "public_strategy_mode",
        "public_strategy_reason",
        "friend_strategy_score",
        "warning_flags",
        "recommendation_confidence",
        "manual_review_flag",
        "close_alternatives",
        "ev_gap_to_second",
        "ev_gap_to_third",
        "decision_note",
        "market_fair_btts_yes_probability",
        "market_fair_btts_no_probability",
        "model_implied_btts_yes_probability",
        "model_implied_btts_no_probability",
        "btts_fit_error",
        "btts_market_available",
        "model_probability_team_a_clean_sheet",
        "model_probability_team_b_clean_sheet",
        "model_probability_no_btts",
        "model_probability_btts",
        "expected_total_goals",
        "expected_team_a_goals",
        "expected_team_b_goals",
        "probability_total_goals_0",
        "probability_total_goals_1",
        "probability_total_goals_2",
        "probability_total_goals_3",
        "probability_total_goals_4_plus",
        "top_5_ev_decomposition",
        "top_5_ev_decomposition_json",
        "top_10_ev_decomposition",
        "top_10_ev_decomposition_json",
        "ev_explanation",
        "extreme_favourite_audit_triggered",
        "normal_grid_recommendation",
        "normal_grid_poisson_recommendation",
        "larger_grid_recommendation",
        "larger_grid_poisson_recommendation",
        "recommendation_changes_with_larger_grid",
        "larger_grid_poisson_changes_recommendation",
        "larger_grid_differs_from_live_recommendation",
        "normal_grid_tail_mass",
        "larger_grid_tail_mass",
        "ev_favourite_3_0_normal_grid",
        "ev_favourite_4_0_normal_grid",
        "ev_favourite_5_0_normal_grid",
        "ev_favourite_3_0_larger_grid",
        "ev_favourite_4_0_larger_grid",
        "ev_favourite_5_0_larger_grid",
    }.issubset(workflow.match_report.columns)
    assert workflow.match_report["recommended_score"].equals(workflow.match_report["final_live_recommended_score"])
    assert "market_consistent_matrix" in workflow.challenger_score_matrices["WC001"]
    assert workflow.match_report["top_5_ev_predictions"].str.len().gt(0).all()
    assert workflow.match_report["plausible_top_alternatives"].str.len().gt(0).all()
    assert not workflow.margin_method_comparison.empty
    assert {"normalised_inverse_odds", "power", "additive", "shin"}.issubset(
        set(workflow.margin_method_comparison["method"])
    )
    assert workflow.match_report["baseline_poisson_recommended_score"].equals(
        workflow.match_report["dixon_coles_recommended_score"]
    )
    assert len(pd.read_csv(settings.csv_output_path)) == 2
    assert len(pd.read_excel(settings.xlsx_output_path)) == 2


def test_btts_audit_uses_aggregated_market_fair_probability(tmp_path: Path) -> None:
    report = run_world_cup_predictions(_settings(tmp_path)).match_report.set_index("match_id")
    yes_one = (1 / 1.95) / ((1 / 1.95) + (1 / 1.87))
    yes_two = (1 / 1.98) / ((1 / 1.98) + (1 / 1.85))
    expected_yes = (yes_one + yes_two) / 2

    assert report.loc["WC001", "market_fair_btts_yes_probability"] == pytest.approx(expected_yes)
    assert report.loc["WC001", "market_fair_btts_no_probability"] == pytest.approx(1 - expected_yes)
    assert report.loc["WC001", "btts_market_available"] == "yes"


def test_model_implied_btts_audit_is_computed_from_score_matrix(tmp_path: Path) -> None:
    workflow = run_world_cup_predictions(_settings(tmp_path))
    report = workflow.match_report.set_index("match_id")
    matrix = workflow.score_matrices["WC001"]

    assert report.loc["WC001", "model_implied_btts_yes_probability"] == pytest.approx(
        matrix.btts_yes_probability()
    )
    assert report.loc["WC001", "model_implied_btts_no_probability"] == pytest.approx(
        1 - matrix.btts_yes_probability()
    )
    assert report.loc["WC001", "model_probability_btts"] == pytest.approx(matrix.btts_yes_probability())
    assert report.loc["WC001", "model_probability_no_btts"] == pytest.approx(1 - matrix.btts_yes_probability())


def test_ev_decomposition_components_sum_to_total_ev(tmp_path: Path) -> None:
    report = run_world_cup_predictions(_settings(tmp_path)).match_report.set_index("match_id")
    records = json.loads(report.loc["WC001", "top_10_ev_decomposition_json"])

    assert len(records) == 10
    for record in records:
        component_total = (
            record["participation_component"]
            + record["result_component"]
            + record["goal_difference_component"]
            + record["exact_score_component"]
        )
        assert component_total == pytest.approx(record["total_expected_points"])
        assert record["total_expected_points_from_components"] == pytest.approx(record["total_expected_points"])


def test_ev_explanation_uses_probability_weighted_margin_wording() -> None:
    explanation = _ev_explanation(
        [
            {
                "predicted_score": "3-0",
                "result_component": 3.6,
                "goal_difference_component": 0.50,
                "exact_score_component": 0.20,
                "participation_component": 1.0,
            },
            {
                "predicted_score": "4-0",
                "result_component": 3.6,
                "goal_difference_component": 0.40,
                "exact_score_component": 0.20,
                "participation_component": 1.0,
            },
        ]
    )

    assert "expected value to the +3 margin than the +4 margin" in explanation
    assert "payoff is higher" not in explanation


def test_decision_aid_flags_clustered_recommendations_for_manual_review() -> None:
    aid = _decision_aid(
        [
            {"predicted_score": "3-0", "total_expected_points": 5.00},
            {"predicted_score": "4-0", "total_expected_points": 4.96},
            {"predicted_score": "5-0", "total_expected_points": 4.94},
        ],
        high_score_cluster=True,
        high_score_cluster_alternatives="3-0, 4-0, 5-0",
    )

    assert aid["recommendation_confidence"] == "low / clustered"
    assert aid["manual_review_flag"] == "yes"
    assert aid["ev_gap_to_second"] == pytest.approx(0.04)
    assert aid["ev_gap_to_third"] == pytest.approx(0.06)
    assert aid["close_alternatives"] == "4-0, 5-0, 3-0"
    assert "High-score cluster" in aid["decision_note"]


def test_decision_aid_marks_clear_ev_lead_as_high_confidence() -> None:
    aid = _decision_aid(
        [
            {"predicted_score": "2-0", "total_expected_points": 5.00},
            {"predicted_score": "1-0", "total_expected_points": 4.80},
            {"predicted_score": "3-0", "total_expected_points": 4.70},
        ],
        high_score_cluster=False,
        high_score_cluster_alternatives="",
    )

    assert aid["recommendation_confidence"] == "high"
    assert aid["manual_review_flag"] == "no"
    assert aid["close_alternatives"] == ""
    assert aid["ev_gap_to_second"] == pytest.approx(0.20)
    assert aid["decision_note"] == "No manual review signal."


def test_plausible_alternatives_suppress_tiny_probability_high_score_duplicates() -> None:
    records = [
        {"predicted_score": "1-0", "total_expected_points": 5.00, "exact_score_probability": 0.08},
        {"predicted_score": "2-1", "total_expected_points": 4.98, "exact_score_probability": 0.05},
        {"predicted_score": "3-2", "total_expected_points": 4.96, "exact_score_probability": 0.012},
        {"predicted_score": "4-3", "total_expected_points": 4.94, "exact_score_probability": 0.004},
        {"predicted_score": "5-4", "total_expected_points": 4.92, "exact_score_probability": 0.0002},
    ]

    layer = _plausible_alternative_decision_layer(
        records,
        config=ProjectConfig(),
        extreme_favourite=False,
        favourite_is_team_a=True,
    )

    assert "1-0" in layer["plausible_top_alternatives"]
    assert "5-4" not in layer["plausible_top_alternatives"]
    assert "5-4" in layer["suppressed_ev_candidates"]
    assert "lower-probability duplicate result/margin bucket" in layer["suppression_reason"]
    assert len(records) == 5


def test_plausible_alternatives_keep_clean_sheet_scores_for_extreme_favourites() -> None:
    records = [
        {"predicted_score": "3-0", "total_expected_points": 5.00, "exact_score_probability": 0.003},
        {"predicted_score": "4-0", "total_expected_points": 4.99, "exact_score_probability": 0.002},
        {"predicted_score": "5-0", "total_expected_points": 4.98, "exact_score_probability": 0.001},
        {"predicted_score": "6-0", "total_expected_points": 4.97, "exact_score_probability": 0.0005},
    ]

    layer = _plausible_alternative_decision_layer(
        records,
        config=ProjectConfig(),
        extreme_favourite=True,
        favourite_is_team_a=True,
    )

    assert "3-0" in layer["plausible_top_alternatives"]
    assert "4-0" in layer["plausible_top_alternatives"]
    assert "5-0" in layer["plausible_top_alternatives"]
    assert "6-0" in layer["plausible_top_alternatives"]
    assert layer["suppressed_ev_candidates"] == ""


def _sample_ev_modal_recommendation() -> GroupPredictionRecommendation:
    return GroupPredictionRecommendation(
        best=GroupPredictionEvaluation((1, 0), 4.20, 0.12, 0.24, 0.42, 0.58),
        alternatives=(
            GroupPredictionEvaluation((1, 1), 3.80, 0.11, 0.26, 0.30, 0.70),
        ),
        most_likely_scoreline=(1, 1),
        differs_from_most_likely=True,
    )


def test_modal_draw_challenger_triggers_for_balanced_modal_draw_setup() -> None:
    diagnostic = _modal_draw_challenger_diagnostic(
        recommendation=_sample_ev_modal_recommendation(),
        favourite_probability=0.42,
        draw_decisive={"draw_vs_decisive_gap": -0.25},
        ev_minus_modal_expected_gap=0.40,
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["modal_draw_challenger_flag"] == "yes"
    assert diagnostic["modal_draw_challenger_score"] == "1-1"
    assert diagnostic["ev_score"] == "1-0"
    assert diagnostic["modal_score"] == "1-1"


def test_modal_draw_challenger_does_not_trigger_for_strong_favourite() -> None:
    diagnostic = _modal_draw_challenger_diagnostic(
        recommendation=_sample_ev_modal_recommendation(),
        favourite_probability=0.78,
        draw_decisive={"draw_vs_decisive_gap": -0.25},
        ev_minus_modal_expected_gap=0.40,
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["modal_draw_challenger_flag"] == "no"
    assert diagnostic["modal_draw_challenger_score"] == ""


def test_modal_draw_challenger_does_not_override_recommended_score() -> None:
    recommendation = _sample_ev_modal_recommendation()
    diagnostic = _modal_draw_challenger_diagnostic(
        recommendation=recommendation,
        favourite_probability=0.42,
        draw_decisive={"draw_vs_decisive_gap": -0.25},
        ev_minus_modal_expected_gap=0.40,
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["modal_draw_challenger_flag"] == "yes"
    assert diagnostic["modal_draw_challenger_score"] != f"{recommendation.best.predicted_score[0]}-{recommendation.best.predicted_score[1]}"
    assert recommendation.best.predicted_score == (1, 0)


def test_draw_prone_flag_triggers_for_low_total_balanced_match() -> None:
    diagnostic = _draw_prone_diagnostic(
        ou_median_total=2.0,
        expected_total_goals=2.1,
        favourite_probability=0.46,
        market_draw_probability=0.31,
        recommendation=_sample_ev_modal_recommendation(),
        draw_decisive={"draw_vs_decisive_gap": -0.20},
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["draw_prone_flag"] == "yes"
    assert "Draw-prone market profile" in diagnostic["draw_prone_reason"]
    assert "expected_total_goals" in diagnostic["total_signal_used_for_draw_prone"]


def test_draw_prone_flag_uses_expected_total_not_ou_median() -> None:
    diagnostic = _draw_prone_diagnostic(
        ou_median_total=3.5,
        expected_total_goals=2.1,
        favourite_probability=0.46,
        market_draw_probability=0.31,
        recommendation=_sample_ev_modal_recommendation(),
        draw_decisive={"draw_vs_decisive_gap": -0.20},
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["draw_prone_flag"] == "yes"
    assert diagnostic["total_signal_used_for_draw_prone"] == "expected_total_goals=2.100"


def test_draw_prone_flag_does_not_trigger_for_strong_favourite() -> None:
    diagnostic = _draw_prone_diagnostic(
        ou_median_total=2.0,
        expected_total_goals=2.1,
        favourite_probability=0.72,
        market_draw_probability=0.20,
        recommendation=_sample_ev_modal_recommendation(),
        draw_decisive={"draw_vs_decisive_gap": -0.20},
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["draw_prone_flag"] == "no"


def test_broad_btts_probability_alone_does_not_trigger_conflict_flag() -> None:
    evaluations = (
        GroupPredictionEvaluation((1, 0), 4.20, 0.12, 0.24, 0.42, 0.58),
        GroupPredictionEvaluation((1, 1), 4.05, 0.11, 0.26, 0.30, 0.70),
    )

    diagnostic = _btts_conflict_diagnostic(
        recommended_score=(1, 0),
        recommended_expected_points=4.20,
        market_btts_yes=0.49,
        model_btts_yes=0.47,
        favourite_probability=0.70,
        draw_prone_flag=False,
        evaluations=evaluations,
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["btts_conflict_flag"] == "no"


def test_btts_conflict_flag_triggers_for_close_btts_alternative() -> None:
    evaluations = (
        GroupPredictionEvaluation((1, 0), 4.20, 0.12, 0.24, 0.42, 0.58),
        GroupPredictionEvaluation((1, 1), 4.05, 0.11, 0.26, 0.30, 0.70),
    )

    diagnostic = _btts_conflict_diagnostic(
        recommended_score=(1, 0),
        recommended_expected_points=4.20,
        market_btts_yes=0.49,
        model_btts_yes=0.47,
        favourite_probability=0.48,
        draw_prone_flag=False,
        evaluations=evaluations,
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["btts_conflict_flag"] == "yes"
    assert diagnostic["best_btts_alternative_score"] == "1-1"
    assert "No-BTTS EV pick conflicts" in diagnostic["btts_conflict_note"]


def test_blowout_risk_flag_triggers_for_strong_favourite_high_total() -> None:
    diagnostic = _blowout_risk_diagnostic(
        favourite_probability=0.74,
        ou_median_total=2.75,
        high_margin_alternatives="2-0 (4.800); 3-0 (4.700)",
        high_score_tail_mass=0.012,
        larger_grid_recommendation="3-0",
        current_recommendation="2-0",
        config=ProjectConfig(enable_margin_method_comparison=False),
    )

    assert diagnostic["blowout_risk_flag"] == "yes"
    assert diagnostic["larger_grid_recommended"] == "yes"
    assert "Blowout-risk profile" in diagnostic["blowout_risk_reason"]


def _dashboard_match_report(**overrides) -> pd.DataFrame:
    row = {
        "match_id": "M001",
        "date": "2026-06-11",
        "time": "18:00",
        "team_a": "Alpha",
        "team_b": "Beta",
        "recommended_score": "1-0",
        "market_consistent_recommended_score": "1-0",
        "dixon_coles_recommended_score": "1-0",
        "correct_score_blended_recommended_score": "1-0",
        "public_strategy_score": "1-0",
        "public_strategy_ev_cost": 0.0,
        "ev_gap_to_second": 0.18,
        "ev_gap_to_third": 0.25,
        "plausible_top_alternatives": "2-0 (4.820); 2-1 (4.700)",
        "high_score_cluster": "no",
        "warning_flags": "",
    }
    row.update(overrides)
    return pd.DataFrame([row])


def _dashboard_margin_rows(score: str = "1-0", differs: str = "no", default_score: str = "1-0") -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "match_id": "M001",
                "team_a": "Alpha",
                "team_b": "Beta",
                "method": "normalised_inverse_odds",
                "method_status": "ok",
                "recommended_score": default_score,
                "differs_from_default": "no",
            },
            {
                "match_id": "M001",
                "team_a": "Alpha",
                "team_b": "Beta",
                "method": "power",
                "method_status": "ok",
                "recommended_score": score,
                "differs_from_default": differs,
            },
        ]
    )


def test_final_decision_dashboard_marks_strong_consensus_high_confidence() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(),
        _dashboard_margin_rows(),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["final_decision_score"] == "1-0"
    assert dashboard["model_consensus"] == "strong_consensus"
    assert dashboard["confidence_level"] == "high"
    assert dashboard["manual_review_flag"] == "no"
    assert dashboard["decision_note"] == "Default recommendation supported by all challenger diagnostics."


def test_final_decision_dashboard_flags_market_consistent_disagreement() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(market_consistent_recommended_score="1-1", ev_gap_to_second=0.03),
        _dashboard_margin_rows(),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["model_consensus"] == "moderate_consensus"
    assert dashboard["confidence_level"] == "low"
    assert dashboard["manual_review_flag"] == "yes"
    assert dashboard["main_alternative_score"] == "1-1"
    assert "market-consistent challenger differs" in dashboard["risk_notes"]
    assert "consider 1-1" in dashboard["decision_note"]


def test_final_decision_dashboard_flags_margin_method_sensitivity() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(),
        _dashboard_margin_rows(score="1-1", differs="yes"),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["margin_method_sensitive"] == "yes"
    assert dashboard["manual_review_flag"] == "yes"
    assert "margin-removal methods change recommendation" in dashboard["risk_notes"]


def test_final_decision_dashboard_flags_market_consistent_failure() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(warning_flags="market_consistent_optimisation_failed"),
        _dashboard_margin_rows(),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["manual_review_flag"] == "yes"
    assert "market_consistent_optimisation_failed" in dashboard["risk_notes"]


def test_market_consistent_failure_warning_propagates_to_report_and_dashboard(monkeypatch) -> None:
    odds = pd.DataFrame(
        [
            {
                "match_id": "MFAIL",
                "date": "2026-06-11",
                "stage": "group stage",
                "group": "A",
                "team_a": "Alpha",
                "team_b": "Beta",
                "bookmaker": "Book",
                "odds_a_win": 2.0,
                "odds_draw": 3.5,
                "odds_b_win": 4.0,
            }
        ]
    )

    def failed_fit(prior, market_probabilities, total_goals=None, correct_score_matrix=None, weights=None):
        return type(
            "FailedMarketConsistentResult",
            (),
            {
                "matrix": prior,
                "diagnostics": {
                    "market_consistent_status": "failed_severe",
                    "market_consistent_optimisation_classification": "optimisation_failed_severe",
                    "market_consistent_optimisation_success": False,
                    "market_consistent_optimisation_status_code": 1,
                    "market_consistent_optimisation_message": "iteration limit reached",
                    "market_consistent_optimisation_iterations": 1,
                    "market_consistent_final_objective_value": 1.0,
                    "market_consistent_gradient_norm": 0.5,
                    "market_consistent_max_constraint_error": 0.5,
                    "market_consistent_constraint_count": 3,
                    "market_consistent_1x2_constraint_count": 3,
                    "market_consistent_btts_constraint_count": 0,
                    "market_consistent_total_goals_constraint_count": 0,
                    "market_consistent_correct_score_constraint_count": 0,
                    "market_consistent_kl_divergence_vs_prior": 0.0,
                    "market_consistent_1x2_fit_error": 0.0,
                    "market_consistent_btts_fit_error": pd.NA,
                    "market_consistent_total_goals_fit_error": pd.NA,
                    "market_consistent_correct_score_fit_error": pd.NA,
                    "market_consistent_asian_totals_used": "",
                },
            },
        )()

    monkeypatch.setattr(workflow_module, "fit_market_consistent_matrix", failed_fit)

    workflow = run_prediction_workflow(
        odds,
        config=ProjectConfig(enable_market_consistent_challenger=True, enable_margin_method_comparison=False),
    )
    report = workflow.match_report.iloc[0]
    dashboard = workflow.final_decision_dashboard.iloc[0]

    assert "market_consistent_optimisation_failed" in report["warning_flags"]
    assert "market_consistent_optimisation_failed_severe" in report["warning_flags"]
    assert report["market_consistent_status"] == "failed_severe"
    assert "Market-consistent optimiser did not converge" in report["warnings"]
    assert "market_consistent_optimisation_failed" in dashboard["warning_flags"]
    assert "market_consistent_optimisation_failed" in dashboard["risk_notes"]


def test_unverified_knockout_scoring_mode_adds_workflow_warning() -> None:
    odds = pd.DataFrame(
        [
            {
                "match_id": "K001",
                "date": "2026-07-01",
                "stage": "round of 16",
                "group": "",
                "team_a": "Alpha",
                "team_b": "Beta",
                "bookmaker": "Book",
                "odds_a_win": 2.0,
                "odds_draw": 3.2,
                "odds_b_win": 4.0,
                "odds_a_qualifies": 1.70,
                "odds_b_qualifies": 2.20,
            }
        ]
    )

    workflow = run_prediction_workflow(
        odds,
        config=ProjectConfig(enable_margin_method_comparison=False, enable_market_consistent_challenger=False),
    )
    report = workflow.match_report.iloc[0]

    assert report["knockout_scoring_mode"] == "unverified"
    assert "knockout_scoring_unverified" in report["warning_flags"]
    assert "Knockout scoring mode is unverified" in report["warnings"]


def test_final_decision_dashboard_marks_margin_method_not_run_when_disabled() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(),
        pd.DataFrame(),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["margin_method_sensitive"] == "not_run"


def test_market_consistent_only_if_close_runs_only_for_risky_matches() -> None:
    config = ProjectConfig(enable_market_consistent_challenger="only_if_close")

    assert not _should_run_market_consistent_challenger(
        setting="only_if_close",
        ev_gap_to_second=0.20,
        extreme_favourite=False,
        warning_flags="",
        config=config,
    )
    assert _should_run_market_consistent_challenger(
        setting="only_if_close",
        ev_gap_to_second=0.03,
        extreme_favourite=False,
        warning_flags="",
        config=config,
    )
    assert _should_run_market_consistent_challenger(
        setting="only_if_close",
        ev_gap_to_second=0.20,
        extreme_favourite=True,
        warning_flags="",
        config=config,
    )


def test_final_decision_dashboard_flags_extreme_favourite_cluster() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(
            recommended_score="3-0",
            market_consistent_recommended_score="3-0",
            dixon_coles_recommended_score="3-0",
            correct_score_blended_recommended_score="3-0",
            public_strategy_score="4-0",
            public_strategy_ev_cost=0.02,
            plausible_top_alternatives="4-0 (4.960); 5-0 (4.800)",
            high_score_cluster="yes",
        ),
        _dashboard_margin_rows(score="3-0", differs="no", default_score="3-0"),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["confidence_level"] == "clustered"
    assert dashboard["high_score_cluster"] == "yes"
    assert dashboard["manual_review_flag"] == "yes"
    assert dashboard["main_alternative_score"] == "4-0"
    assert "High-score cluster" in dashboard["decision_note"]


def test_final_decision_dashboard_flags_modal_draw_challenger_for_manual_review() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(
            modal_draw_challenger_flag="yes",
            modal_draw_challenger_score="1-1",
            btts_conflict_flag="no",
        ),
        _dashboard_margin_rows(),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["final_decision_score"] == "1-0"
    assert dashboard["modal_draw_challenger_flag"] == "yes"
    assert dashboard["manual_review_flag"] == "yes"
    assert "WC 2022 backtest" in dashboard["decision_note"]


def test_pattern_flags_do_not_override_final_decision_score() -> None:
    dashboard = _build_final_decision_dashboard(
        _dashboard_match_report(
            recommended_score="1-0",
            draw_prone_flag="yes",
            draw_prone_reason="Draw-prone market profile",
            blowout_risk_flag="no",
        ),
        _dashboard_margin_rows(),
        ProjectConfig(enable_margin_method_comparison=False),
    ).iloc[0]

    assert dashboard["final_decision_score"] == "1-0"
    assert dashboard["manual_review_flag"] == "yes"
    assert "Draw-prone market profile" in dashboard["decision_note"]


def test_btts_warning_appears_for_recommendation_conflicting_with_strong_market_signal() -> None:
    flags = _btts_warning_flags(
        has_btts=True,
        market_btts_yes=0.60,
        market_btts_no=0.40,
        model_btts_yes=0.58,
        recommended_score=(1, 0),
    )

    assert "recommended_no_btts_against_strong_btts_yes_market" in flags


def test_top_ten_ev_decomposition_is_written_to_excel(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    run_world_cup_predictions(settings)
    workbook = load_workbook(settings.xlsx_output_path)
    recommendation_headers = [cell.value for cell in workbook["recommendations"][1]]
    diagnostics_headers = [cell.value for cell in workbook["ev_decomposition"][1]]

    assert "top_10_ev_decomposition" in recommendation_headers
    assert "recommendation_confidence" in recommendation_headers
    assert "manual_review_flag" in recommendation_headers
    assert "decision_note" in recommendation_headers
    assert "plausible_top_alternatives" in recommendation_headers
    assert "suppressed_ev_candidates" in recommendation_headers
    assert "suppression_reason" in recommendation_headers
    assert "ev_decomposition" in workbook.sheetnames
    assert "market_consistent" in workbook.sheetnames
    assert "margin_methods" in workbook.sheetnames
    assert "final_decision_dashboard" in workbook.sheetnames
    assert "predicted_score" in diagnostics_headers
    assert "exact_score_component" in diagnostics_headers
    market_consistent_headers = [cell.value for cell in workbook["market_consistent"][1]]
    dashboard_headers = [cell.value for cell in workbook["final_decision_dashboard"][1]]
    assert "market_consistent_top_10_ev_scorelines" in recommendation_headers
    assert "market_consistent_status" in market_consistent_headers
    assert "market_consistent_optimisation_classification" in market_consistent_headers
    assert "market_consistent_optimisation_message" in market_consistent_headers
    assert "market_consistent_gradient_norm" in market_consistent_headers
    assert "market_consistent_max_constraint_error" in market_consistent_headers
    assert "market_consistent_constraint_count" in market_consistent_headers
    assert "warning_flags" in market_consistent_headers
    assert "market_consistent_top_10_probability_scorelines" in market_consistent_headers
    assert len(market_consistent_headers) == len(set(market_consistent_headers))
    assert "final_decision_score" in recommendation_headers
    assert "final_decision_score" in dashboard_headers
    assert "risk_notes" in dashboard_headers
    assert workbook["ev_decomposition"].max_row == 21


def test_margin_method_comparison_does_not_change_default_recommended_score(tmp_path: Path) -> None:
    workflow = run_world_cup_predictions(_settings(tmp_path))
    report = workflow.match_report.set_index("match_id")
    comparison = workflow.margin_method_comparison
    default_rows = comparison[comparison["method"] == "normalised_inverse_odds"].set_index("match_id")

    assert report["recommended_score"].equals(report["final_live_recommended_score"])
    assert default_rows["recommended_score"].equals(report.loc[default_rows.index, "recommended_score"])
    assert default_rows["differs_from_default"].eq("no").all()


def test_configured_margin_method_changes_active_workflow_probabilities(tmp_path: Path) -> None:
    default = run_world_cup_predictions(
        _settings(tmp_path / "default"),
        ProjectConfig(enable_margin_method_comparison=False),
    ).match_report.set_index("match_id")
    power = run_world_cup_predictions(
        _settings(tmp_path / "power"),
        ProjectConfig(margin_removal_method="power", enable_margin_method_comparison=False),
    ).match_report.set_index("match_id")

    assert not default["market_a_win"].equals(power["market_a_win"])
    assert power["margin_removal_method"].eq("power").all()


def test_configured_shin_margin_method_is_used_by_active_workflow(tmp_path: Path) -> None:
    report = run_world_cup_predictions(
        _settings(tmp_path),
        ProjectConfig(margin_removal_method="shin", enable_margin_method_comparison=False),
    ).match_report

    assert report["margin_removal_method"].eq("shin").all()
    assert report["recommended_score"].equals(report["final_live_recommended_score"])


def test_extreme_favourite_audit_and_larger_grid_sensitivity_trigger(tmp_path: Path) -> None:
    input_path = tmp_path / "extreme.csv"
    pd.DataFrame(
        [
            {
                "match_id": "EXTREME",
                "date": "2026-06-12",
                "stage": "group stage",
                "group": "A",
                "team_a": "Favourite",
                "team_b": "Longshot",
                "bookmaker": "MarketOne",
                "odds_a_win": 1.08,
                "odds_draw": 12.0,
                "odds_b_win": 35.0,
                "odds_btts_yes": 2.30,
                "odds_btts_no": 1.65,
            }
        ]
    ).to_csv(input_path, index=False)
    settings = WorldCupPredictionSettings(
        input_path=input_path,
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    report = run_world_cup_predictions(settings).match_report.iloc[0]

    assert report["favourite_probability"] > 0.75
    assert bool(report["extreme_favourite_audit_triggered"])
    assert report["top_clean_sheet_scores"]
    assert report["top_favourite_margin_probabilities"]
    assert report["current_final_recommendation"] == report["recommended_score"]
    assert report["normal_grid_poisson_recommendation"] == report["baseline_poisson_recommended_score"]
    assert report["larger_grid_recommendation"]
    assert report["larger_grid_poisson_recommendation"] == report["larger_grid_recommendation"]
    assert pd.notna(report["larger_grid_tail_mass"])
    assert report["larger_grid_tail_mass"] < report["normal_grid_tail_mass"]
    assert pd.notna(report["ev_favourite_3_0_normal_grid"])
    assert pd.notna(report["ev_favourite_3_0_larger_grid"])


def _extreme_favourite_settings(tmp_path: Path) -> WorldCupPredictionSettings:
    input_path = tmp_path / "extreme_dynamic.csv"
    pd.DataFrame(
        [
            {
                "match_id": "EXTREME",
                "date": "2026-06-12",
                "stage": "group stage",
                "group": "A",
                "team_a": "Favourite",
                "team_b": "Longshot",
                "bookmaker": "MarketOne",
                "odds_a_win": 1.06,
                "odds_draw": 14.0,
                "odds_b_win": 41.0,
                "odds_btts_yes": 2.40,
                "odds_btts_no": 1.60,
            }
        ]
    ).to_csv(input_path, index=False)
    return WorldCupPredictionSettings(
        input_path=input_path,
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )


def test_dynamic_grid_uses_extreme_favourite_max_goals(tmp_path: Path) -> None:
    settings = _extreme_favourite_settings(tmp_path)

    report = run_world_cup_predictions(
        settings, ProjectConfig(extreme_favourite_max_goals=12)
    ).match_report.iloc[0]

    assert bool(report["extreme_favourite_audit_triggered"])
    assert report["grid_max_goals_used"] == 12
    assert report["tail_mass_after_grid_extension"] < report["tail_mass_before_grid_extension"]
    assert report["recommendation_changed_due_to_larger_grid"] in {"yes", "no"}


def test_dynamic_grid_can_be_disabled(tmp_path: Path) -> None:
    settings = _extreme_favourite_settings(tmp_path)

    report = run_world_cup_predictions(
        settings, ProjectConfig(dynamic_grid_enabled=False, max_goals_score_matrix=8)
    ).match_report.iloc[0]

    # The diagnostic grid extension is suppressed; the default grid is reported.
    assert report["grid_max_goals_used"] == 8
    assert report["recommendation_changed_due_to_larger_grid"] == "no"


def test_normal_match_reports_default_grid(tmp_path: Path) -> None:
    settings = WorldCupPredictionSettings(
        input_path=EXAMPLES / "example_world_cup_odds.csv",
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    report = run_world_cup_predictions(settings, ProjectConfig(max_goals_score_matrix=8)).match_report

    balanced = report[~report["extreme_favourite_audit_triggered"].astype(bool)]
    assert not balanced.empty
    assert (balanced["grid_max_goals_used"] == 8).all()


def test_larger_grid_sensitivity_can_change_poisson_recommendation_on_synthetic_extreme_favourite(
    tmp_path: Path,
) -> None:
    input_path = tmp_path / "extreme_small_grid.csv"
    pd.DataFrame(
        [
            {
                "match_id": "EXTREME_GRID",
                "date": "2026-06-12",
                "stage": "group stage",
                "group": "A",
                "team_a": "Favourite",
                "team_b": "Longshot",
                "bookmaker": "MarketOne",
                "odds_a_win": 1.04,
                "odds_draw": 17.0,
                "odds_b_win": 55.0,
                "odds_btts_yes": 2.60,
                "odds_btts_no": 1.55,
            }
        ]
    ).to_csv(input_path, index=False)
    settings = WorldCupPredictionSettings(
        input_path=input_path,
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    report = run_world_cup_predictions(
        settings,
        ProjectConfig(max_goals_score_matrix=2, max_candidate_goals=5),
    ).match_report.iloc[0]

    assert bool(report["extreme_favourite_audit_triggered"])
    assert bool(report["larger_grid_poisson_changes_recommendation"])
    assert report["normal_grid_poisson_recommendation"] != report["larger_grid_poisson_recommendation"]
    assert report["larger_grid_tail_mass"] < report["normal_grid_tail_mass"]
    assert "larger_grid_sensitivity" in report["warning_flags"]


def test_world_cup_workflow_optionally_uses_correct_score_blend(tmp_path: Path) -> None:
    settings = WorldCupPredictionSettings(
        input_path=EXAMPLES / "example_world_cup_odds.csv",
        correct_score_input_path=EXAMPLES / "example_world_cup_correct_score_odds.csv",
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    workflow = run_world_cup_predictions(
        settings,
        ProjectConfig(correct_score_poisson_weight=0.5, min_scorelines_for_blend=6),
    )
    report = workflow.match_report.set_index("match_id")

    assert report["has_correct_score_market"].all()
    assert (report["correct_score_poisson_weight"] == 0.5).all()
    assert report["correct_score_market_top_10"].str.len().gt(0).all()
    assert report["correct_score_blended_top_10"].str.len().gt(0).all()
    assert report["correct_score_kl_divergence"].gt(0).all()
    assert (report["correct_score_scorelines_count"] == 6).all()
    assert (report["correct_score_bookmakers_count"] == 2).all()
    assert report["correct_score_sparse_warning"].str.contains("fewer_than_10_scorelines").all()
    assert report["top_10_market_scorelines"].equals(report["correct_score_market_top_10"])
    assert report["top_10_blended_scorelines"].equals(report["correct_score_blended_top_10"])
    assert report["kl_divergence"].equals(report["correct_score_kl_divergence"])
    assert (report["selected_blend_weight"] == 0.5).all()
    assert report["correct_score_blend_applied"].all()
    assert (report["correct_score_aggregation_method"] == "mean").all()
    assert (report["number_of_correct_score_bookmakers"] == 2).all()
    assert report["average_correct_score_overround"].gt(0).all()
    assert report["max_correct_score_overround"].gt(0).all()
    assert report["outlier_count"].eq(0).all()
    assert report["scoreline_coverage_warning"].str.contains("missing_common_scorelines").all()
    assert report["correct_score_bookmaker_diagnostics"].str.contains("overround=").all()
    assert report["dixon_coles_rho_source"].eq("market_estimated").all()
    assert report["dixon_coles_rho_used"].between(-0.20, 0.20).all()


def test_market_specific_devig_keeps_correct_score_conservative(tmp_path: Path) -> None:
    from wc_predictor.config import DevigConfig

    settings = WorldCupPredictionSettings(
        input_path=EXAMPLES / "example_world_cup_odds.csv",
        correct_score_input_path=EXAMPLES / "example_world_cup_correct_score_odds.csv",
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    workflow = run_world_cup_predictions(
        settings,
        ProjectConfig(
            correct_score_poisson_weight=0.5,
            min_scorelines_for_blend=6,
            devig=DevigConfig(one_x_two_method="power", correct_score_method="normalised_inverse_odds"),
        ),
    )
    report = workflow.match_report

    # 1X2 used the aggressive method while correct score stayed conservative,
    # and the run did not crash on the many-outcome correct-score market.
    assert report["devig_methods_by_market"].str.contains("1x2:power->power").all()
    assert report["devig_methods_by_market"].str.contains(
        "correct_score:normalised_inverse_odds"
    ).all()


def test_market_specific_devig_does_not_crash_with_shin_on_correct_score(tmp_path: Path) -> None:
    from wc_predictor.config import DevigConfig

    settings = WorldCupPredictionSettings(
        input_path=EXAMPLES / "example_world_cup_odds.csv",
        correct_score_input_path=EXAMPLES / "example_world_cup_correct_score_odds.csv",
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    # Even an aggressive correct-score method must fall back safely, never crash.
    workflow = run_world_cup_predictions(
        settings,
        ProjectConfig(
            correct_score_poisson_weight=0.5,
            min_scorelines_for_blend=6,
            devig=DevigConfig(default_method="shin", correct_score_method="shin"),
        ),
    )
    assert not workflow.match_report.empty


def _market_consistent_settings(tmp_path: Path) -> WorldCupPredictionSettings:
    return WorldCupPredictionSettings(
        input_path=EXAMPLES / "example_world_cup_odds.csv",
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )


def test_market_consistent_default_prior_is_independent_poisson(tmp_path: Path) -> None:
    report = run_world_cup_predictions(
        _market_consistent_settings(tmp_path),
        ProjectConfig(enable_market_consistent_challenger=True, enable_margin_method_comparison=False),
    ).match_report

    assert (report["market_consistent_prior_source"] == "independent_poisson").all()
    assert (report["market_consistent_prior_kl_vs_independent"] == 0.0).all()
    assert (report["market_consistent_prior_changed_recommendation"] == "no").all()


def test_market_consistent_dixon_coles_prior_is_reported(tmp_path: Path) -> None:
    report = run_world_cup_predictions(
        _market_consistent_settings(tmp_path),
        ProjectConfig(
            enable_market_consistent_challenger=True,
            enable_margin_method_comparison=False,
            market_consistent_prior="dixon_coles",
            enable_dixon_coles_rho_estimation=False,
            dixon_coles_rho=0.12,
        ),
    ).match_report

    assert (report["market_consistent_prior_source"] == "dixon_coles").all()
    assert (report["market_consistent_prior_dixon_coles_rho_used"] == 0.12).all()
    assert (report["market_consistent_prior_kl_vs_independent"] > 0).all()


def test_bivariate_poisson_diagnostic_reported_when_enabled(tmp_path: Path) -> None:
    report = run_world_cup_predictions(
        _market_consistent_settings(tmp_path),
        ProjectConfig(
            enable_bivariate_poisson_diagnostic=True,
            bivariate_poisson_covariance=0.2,
            enable_margin_method_comparison=False,
        ),
    ).match_report

    assert report["bivariate_prior_recommended_score"].str.len().gt(0).all()
    assert (report["bivariate_lambda_3"] > 0).any()


def test_market_consistent_constraint_correlation_note_is_reported(tmp_path: Path) -> None:
    report = run_world_cup_predictions(
        _market_consistent_settings(tmp_path),
        ProjectConfig(enable_market_consistent_challenger=True, enable_margin_method_comparison=False),
    ).match_report

    assert report["market_consistent_constraint_correlation_note"].str.contains(
        "constraints_treated_independently"
    ).all()
    assert report["market_consistent_active_constraint_groups"].str.contains("1x2").all()


def test_market_consistent_group_weights_can_downweight_a_group(tmp_path: Path) -> None:
    from wc_predictor.config import MarketConsistentGroupWeights

    base = run_world_cup_predictions(
        _market_consistent_settings(tmp_path),
        ProjectConfig(enable_market_consistent_challenger=True, enable_margin_method_comparison=False),
    ).match_report.set_index("match_id")
    downweighted = run_world_cup_predictions(
        _market_consistent_settings(tmp_path),
        ProjectConfig(
            enable_market_consistent_challenger=True,
            enable_margin_method_comparison=False,
            market_consistent_group_weights=MarketConsistentGroupWeights(one_x_two=0.01),
        ),
    ).match_report.set_index("match_id")

    assert not base["market_consistent_kl_divergence_vs_prior"].equals(
        downweighted["market_consistent_kl_divergence_vs_prior"]
    )


def test_bivariate_poisson_diagnostic_disabled_by_default(tmp_path: Path) -> None:
    report = run_world_cup_predictions(
        _market_consistent_settings(tmp_path),
        ProjectConfig(enable_margin_method_comparison=False),
    ).match_report

    assert (report["bivariate_prior_recommended_score"] == "").all()
    assert (report["bivariate_differs_from_default"] == "no").all()


def test_world_cup_workflow_optionally_loads_total_goals_ladder(tmp_path: Path) -> None:
    total_goals_path = tmp_path / "total_goals.csv"
    pd.DataFrame(
        [
            {"match_id": "WC001", "bookmaker": "MarketOne", "line": 1.5, "odds_over": 1.40, "odds_under": 3.00},
            {"match_id": "WC001", "bookmaker": "MarketTwo", "line": 1.5, "odds_over": 1.45, "odds_under": 2.90},
            {"match_id": "WC001", "bookmaker": "MarketOne", "line": 2.5, "odds_over": 2.00, "odds_under": 1.80},
            {"match_id": "WC001", "bookmaker": "MarketTwo", "line": 2.5, "odds_over": 2.10, "odds_under": 1.75},
        ]
    ).to_csv(total_goals_path, index=False)
    baseline = _settings(tmp_path)
    settings = WorldCupPredictionSettings(
        input_path=baseline.input_path,
        total_goals_input_path=total_goals_path,
        csv_output_path=baseline.csv_output_path,
        xlsx_output_path=baseline.xlsx_output_path,
        submission_xlsx_output_path=baseline.submission_xlsx_output_path,
    )

    report = run_world_cup_predictions(settings).match_report.set_index("match_id")

    assert report.loc["WC001", "total_goals_lines_available"] == "1.5; 2.5"
    assert bool(report.loc["WC001", "multi_line_totals_used"])
    assert report.loc["WC002", "total_goals_lines_available"] == ""


def test_world_cup_excel_export_is_formatted_for_manual_review(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    run_world_cup_predictions(settings)
    worksheet = load_workbook(settings.xlsx_output_path)["recommendations"]
    headers = [cell.value for cell in worksheet[1]]

    assert worksheet.freeze_panes == "A2"
    assert headers[:30] == [
        "match_id",
        "date",
        "time",
        "stage",
        "group",
        "team_a",
        "team_b",
        "recommended_score",
        "final_decision_score",
        "confidence_level",
        "manual_review_flag",
        "main_alternative_score",
        "plausible_alternatives",
        "strategic_alternative_score",
        "override_candidate",
        "model_consensus",
        "risk_notes",
        "decision_note",
        "market_a",
        "market_draw",
        "market_b",
        "favourite_probability",
        "favourite_bucket",
        "lambda_a",
        "lambda_b",
        "recommended_qualifier",
        "best_expected_points",
        "most_likely_scoreline",
        "ev_optimal_differs",
        "top_5_ev_predictions",
    ]
    column = {name: index + 1 for index, name in enumerate(headers)}
    assert worksheet.cell(2, column["market_a"]).number_format == "0.00%"
    assert worksheet.cell(2, column["favourite_probability"]).number_format == "0.00%"
    assert worksheet.cell(2, column["lambda_a"]).number_format == "0.0000"
    assert worksheet.cell(2, column["calibration_loss"]).number_format == "0.0000"
    assert worksheet.cell(2, column["best_expected_points"]).number_format == "0.000"
    assert worksheet.cell(2, column["estimated_most_crowded_public_pick_share"]).number_format == "0.00%"
    assert worksheet.cell(2, column["public_strategy_ev_cost"]).number_format == "0.000"
    assert worksheet.cell(2, column["top_5_ev_predictions"]).alignment.wrap_text
    assert worksheet.cell(2, column["public_strategy_reason"]).alignment.wrap_text
    assert worksheet.cell(2, column["warnings"]).alignment.wrap_text
    assert worksheet.column_dimensions["A"].width >= len("match_id")


def test_world_cup_submission_sheet_contains_only_entry_columns_and_formatting(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    run_world_cup_predictions(settings)
    worksheet = load_workbook(settings.submission_xlsx_output_path)["submission"]
    headers = [cell.value for cell in worksheet[1]]
    column = {name: index + 1 for index, name in enumerate(headers)}

    assert headers == [
        "match_id",
        "date",
        "stage",
        "group",
        "team_a",
        "team_b",
        "recommended_score",
        "final_decision_score",
        "confidence_level",
        "manual_review_flag",
        "main_alternative_score",
        "plausible_alternatives",
        "decision_note",
        "recommended_qualifier",
        "best_expected_points",
        "recommendation_confidence",
        "plausible_top_alternatives",
        "suppressed_ev_candidates",
        "close_alternatives",
        "ev_gap_to_second",
        "ev_gap_to_third",
        "favourite_bucket",
        "top_3_alternatives",
        "notes",
    ]
    assert worksheet.freeze_panes == "A2"
    assert worksheet.cell(2, column["best_expected_points"]).number_format == "0.000"
    assert worksheet.cell(2, column["ev_gap_to_second"]).number_format == "0.000"
    assert worksheet.cell(2, column["top_3_alternatives"]).alignment.wrap_text
    assert worksheet.cell(2, column["plausible_alternatives"]).alignment.wrap_text
    assert worksheet.cell(2, column["plausible_top_alternatives"]).alignment.wrap_text
    assert worksheet.cell(2, column["suppressed_ev_candidates"]).alignment.wrap_text
    assert worksheet.cell(2, column["decision_note"]).alignment.wrap_text
    assert worksheet.cell(2, column["notes"]).alignment.wrap_text
    assert worksheet.cell(2, column["top_3_alternatives"]).value.count(";") == 2
    assert worksheet.cell(2, column["match_id"]).value == "WC001"
    assert worksheet.cell(3, column["match_id"]).value == "WC002"
    assert worksheet.column_dimensions["A"].width >= len("match_id")


def test_create_world_cup_odds_file_copies_template(tmp_path: Path) -> None:
    destination = tmp_path / "raw" / "world_cup_odds.xlsx"
    path, changed = create_world_cup_odds_file(destination)

    assert changed
    assert path == destination
    assert destination.read_bytes() == (TEMPLATES / "world_cup_odds_template.xlsx").read_bytes()


def test_create_world_cup_odds_file_does_not_overwrite_existing_file(tmp_path: Path) -> None:
    destination = tmp_path / "world_cup_odds.xlsx"
    destination.write_text("manual odds", encoding="utf-8")

    _, changed = create_world_cup_odds_file(destination)

    assert not changed
    assert destination.read_text(encoding="utf-8") == "manual odds"

    _, changed = create_world_cup_odds_file(destination, overwrite=True)

    assert changed
    assert destination.read_bytes() == (TEMPLATES / "world_cup_odds_template.xlsx").read_bytes()


def test_world_cup_workflow_loads_xlsx_odds_input(tmp_path: Path) -> None:
    input_path = tmp_path / "world_cup_odds.xlsx"
    pd.read_csv(EXAMPLES / "example_world_cup_odds.csv").to_excel(input_path, index=False)
    settings = WorldCupPredictionSettings(
        input_path=input_path,
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    workflow = run_world_cup_predictions(settings)

    assert len(workflow.match_report) == 2
    assert workflow.match_report.loc[0, "match_id"] == "WC001"


def test_optional_odds_source_metadata_is_summarised_when_present(tmp_path: Path) -> None:
    input_path = tmp_path / "world_cup_odds.xlsx"
    odds = pd.read_csv(EXAMPLES / "example_world_cup_odds.csv")
    odds["odds_timestamp"] = [
        "2026-06-01T09:00:00Z",
        "2026-06-01T11:00:00Z",
        "2026-06-02T09:00:00Z",
        "2026-06-02T11:00:00Z",
    ]
    odds["odds_source_url"] = ["https://one.example", "https://two.example"] * 2
    odds["source_quality"] = ["major_bookmaker", "sharp"] * 2
    odds.to_excel(input_path, index=False)
    settings = WorldCupPredictionSettings(
        input_path=input_path,
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    report = run_world_cup_predictions(settings).match_report.set_index("match_id")

    assert report.loc["WC001", "odds_timestamp_min"] == "2026-06-01T09:00:00+00:00"
    assert report.loc["WC001", "odds_timestamp_max"] == "2026-06-01T11:00:00+00:00"
    assert report.loc["WC001", "odds_source_urls"] == "https://one.example; https://two.example"
    assert report.loc["WC001", "source_qualities"] == "major_bookmaker; sharp"
    assert report.loc["WC001", "bookmakers_used"] == "MarketOne; MarketTwo"
    assert report.loc["WC001", "number_of_bookmakers"] == 2


def test_world_cup_warning_flags_cover_model_risk_conditions(tmp_path: Path) -> None:
    input_path = tmp_path / "risky_odds.csv"
    pd.DataFrame(
        [
            {
                "match_id": "RISKY",
                "date": "2026-07-01",
                "stage": "round of 32",
                "group": "",
                "team_a": "Favourite",
                "team_b": "Longshot",
                "bookmaker": "OnlyBook",
                "odds_a_win": 1.03,
                "odds_draw": 20.0,
                "odds_b_win": 50.0,
                "odds_timestamp": "2000-01-01T00:00:00Z",
            }
        ]
    ).to_csv(input_path, index=False)
    settings = WorldCupPredictionSettings(
        input_path=input_path,
        csv_output_path=tmp_path / "recommendations.csv",
        xlsx_output_path=tmp_path / "recommendations.xlsx",
        submission_xlsx_output_path=tmp_path / "submission.xlsx",
    )

    report = run_world_cup_predictions(
        settings,
        ProjectConfig(max_goals_score_matrix=0, poor_calibration_loss_threshold=-1.0),
    ).match_report.iloc[0]
    flags = set(report["warning_flags"].split("; "))

    assert {
        "only_one_bookmaker",
        "no_over_under",
        "no_btts",
        "high_calibration_error",
        "high_tail_mass",
        "stale_odds_timestamp",
        "extreme_favourite",
        "knockout_missing_qualification_odds",
    }.issubset(flags)


def test_world_cup_model_risk_summary_reports_compact_counts(tmp_path: Path) -> None:
    report = run_world_cup_predictions(_settings(tmp_path)).match_report.copy()
    report.loc[0, "warning_flags"] = "only_one_bookmaker; no_over_under; no_btts; extreme_favourite"
    report.loc[1, "warning_flags"] = "high_calibration_error; knockout_missing_qualification_odds"

    summary = format_world_cup_model_risk_summary(report)

    assert "Model-Risk Summary" in summary
    assert "- Matches: 2" in summary
    assert "- Matches with only one bookmaker: 1" in summary
    assert "- Matches without over/under odds: 1" in summary
    assert "- Matches without BTTS odds: 1" in summary
    assert "- Extreme favourites: 1" in summary
    assert "- Matches with high calibration error: 1" in summary
    assert "- Knockout matches missing qualification odds: 1" in summary


def test_model_risk_documentation_exists() -> None:
    assert Path("docs/stylised_facts.md").exists()
    assert Path("docs/model_roadmap.md").exists()


def test_default_world_cup_input_prefers_xlsx_when_both_exist(tmp_path: Path, monkeypatch) -> None:
    prepared = tmp_path / "input" / "prepared"
    prepared.mkdir(parents=True)
    (prepared / "world_cup_odds.csv").touch()
    (prepared / "world_cup_odds.xlsx").touch()
    monkeypatch.chdir(tmp_path)

    assert resolve_world_cup_odds_input() == Path("input/prepared/world_cup_odds.xlsx")


def test_world_cup_console_summary_contains_manual_inspection_fields(tmp_path: Path) -> None:
    settings = _settings(tmp_path)
    workflow = run_world_cup_predictions(settings)
    summary = format_world_cup_console_summary(
        workflow.match_report,
        settings.input_path,
        settings.csv_output_path,
        settings.xlsx_output_path,
        settings.submission_xlsx_output_path,
    )

    assert "World Cup Predictions" in summary
    assert "WC001 | Alpha vs Beta" in summary
    assert "Stage/group: group stage / Group A" in summary
    assert "Fair 1X2:" in summary
    assert "Calibrated lambdas:" in summary
    assert "Favourite strength: p_fav=" in summary
    assert "Modal scoreline:" in summary
    assert "EV-optimal scoreline:" in summary
    assert "Top 5 EV scorelines:" in summary
    assert "Public-ranking strategy:" in summary
    assert "Estimated most crowded public score:" in summary
    assert "Public strategy score:" in summary
    assert "CSV recommendations:" in summary
    assert "Excel recommendations:" in summary
    assert "Submission sheet:" in summary
    assert "Prediction Output Diagnostics" in summary
    assert "Recommended Score Counts" in summary
    assert "Favourite-Strength Bucket Counts" in summary
    assert "- Average lambda_a:" in summary
    assert "- Average lambda_b:" in summary
    assert "- EV-optimal score differs from modal scoreline:" in summary
    assert "Top 10 Matches By Favourite Probability" in summary
    assert "Model-Risk Summary" in summary
    assert "market_a=" in summary
    assert "market_draw=" in summary
    assert "market_b=" in summary


def test_world_cup_prediction_diagnostics_show_only_ten_strongest_favourites(tmp_path: Path) -> None:
    report = run_world_cup_predictions(_settings(tmp_path)).match_report
    rows = []
    for index in range(11):
        row = report.iloc[0].copy()
        row["match_id"] = f"M{index:02d}"
        row["favourite_probability"] = 0.40 + index / 100
        rows.append(row)

    diagnostics = format_world_cup_prediction_diagnostics(pd.DataFrame(rows))

    assert "M10 | Alpha vs Beta | p_fav=50.00%" in diagnostics
    assert "M01 | Alpha vs Beta | p_fav=41.00%" in diagnostics
    assert "M00 | Alpha vs Beta" not in diagnostics
    assert diagnostics.index("M10 | Alpha vs Beta") < diagnostics.index("M09 | Alpha vs Beta")


def test_world_cup_prediction_script_runs_with_overrides(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    settings = _settings(tmp_path)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_world_cup_predictions.py",
            "--odds",
            str(settings.input_path),
            "--csv-output",
            str(settings.csv_output_path),
            "--xlsx-output",
            str(settings.xlsx_output_path),
            "--submission-xlsx-output",
            str(settings.submission_xlsx_output_path),
        ],
    )

    runpy.run_path(str(RUN_SCRIPT), run_name="__main__")

    assert settings.csv_output_path.exists()
    assert settings.xlsx_output_path.exists()
    assert settings.submission_xlsx_output_path.exists()
    assert "WC002 | Gamma vs Delta" in capsys.readouterr().out


def test_world_cup_prediction_script_prints_friendly_missing_file_message(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["run_world_cup_predictions.py"])

    runpy.run_path(str(RUN_SCRIPT), run_name="__main__")

    assert capsys.readouterr().out.strip() == WORLD_CUP_ODDS_MISSING_MESSAGE


def test_create_world_cup_odds_file_script_runs_with_overrides(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    destination = tmp_path / "world_cup_odds.xlsx"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "create_world_cup_odds_file.py",
            "--output",
            str(destination),
            "--template",
            str(TEMPLATES.resolve() / "world_cup_odds_template.xlsx"),
        ],
    )

    runpy.run_path(str(CREATE_SCRIPT), run_name="__main__")

    assert destination.exists()
    assert "Fill in" in capsys.readouterr().out
