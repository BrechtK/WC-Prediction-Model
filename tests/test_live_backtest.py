"""Tests for the live-pipeline historical backtest."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from wc_predictor.live_backtest import (
    DEFAULT_BACKTEST_CONFIGS,
    QUICK_BACKTEST_CONFIGS,
    RESEARCH_BACKTEST_CONFIGS,
    HistoricalResult,
    LiveBacktestSettings,
    _build_round_summary,
    _build_summary,
    _classify_missing_odds,
    _format_expected_vs_actual_summary,
    _format_round_summary,
    _group_stage_playing_round,
    _parse_score,
    _score_strategies,
    load_historical_results,
    run_live_backtest,
)

CORE_FIXTURES = Path("tests/fixtures/oddsportal_core_pastes")


# ---------------------------------------------------------------------------
# Fixtures / helpers
# ---------------------------------------------------------------------------

def _write_results_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_csv(path, index=False)


def _write_combined_paste(folder: Path, match_id: str, *, include_ah: bool = False) -> None:
    """Write a minimal combined paste .txt file in live-tool format."""
    folder.mkdir(parents=True, exist_ok=True)

    one_x_two = CORE_FIXTURES / "MCORE_1x2.txt"
    over_under = CORE_FIXTURES / "MCORE_over_under.txt"
    btts = CORE_FIXTURES / "MCORE_btts.txt"

    sections = [
        ("MATCH", f"Team A vs Team B"),
        ("1X2", one_x_two.read_text(encoding="utf-8")),
        ("OVER_UNDER", over_under.read_text(encoding="utf-8")),
        ("BTTS", btts.read_text(encoding="utf-8")),
    ]
    if include_ah:
        ah_text = "\n".join([
            "Asian Handicap -1.5",
            "Pinnacle",
            "-1.5",
            "1.85",
            "1.97",
            "Bet365",
            "-1.5",
            "1.83",
            "1.95",
        ])
        sections.append(("ASIAN_HANDICAP", ah_text))

    (folder / f"{match_id}.txt").write_text(
        "\n\n".join(f"### {header}\n{content}" for header, content in sections),
        encoding="utf-8",
    )


def _simple_settings(
    tmp_path: Path,
    match_ids: list[str],
    *,
    include_ah: bool = False,
    configs=None,
) -> LiveBacktestSettings:
    odds_folder = tmp_path / "odds"
    results_path = tmp_path / "results.csv"
    cache_folder = tmp_path / "cache"

    for match_id in match_ids:
        _write_combined_paste(odds_folder, match_id, include_ah=include_ah)

    _write_results_csv(
        results_path,
        [
            {
                "match_id": mid,
                "date": "2022-11-20",
                "stage": "group stage",
                "group": "A",
                "team_a": "Team A",
                "team_b": "Team B",
                "actual_score_a": 1,
                "actual_score_b": 0,
            }
            for mid in match_ids
        ],
    )
    from wc_predictor.config import ProjectConfig

    if configs is None:
        configs = (
            DEFAULT_BACKTEST_CONFIGS[0],  # baseline_ev
        )

    return LiveBacktestSettings(
        historical_odds_folder=odds_folder,
        results_path=results_path,
        cache_folder=cache_folder,
        summary_output_path=tmp_path / "summary.csv",
        predictions_output_path=tmp_path / "predictions.csv",
        excel_output_path=tmp_path / "backtest.xlsx",
        configs=configs,
    )


# ---------------------------------------------------------------------------
# Unit tests: load_historical_results
# ---------------------------------------------------------------------------

def test_load_historical_results_parses_required_columns(tmp_path: Path) -> None:
    path = tmp_path / "results.csv"
    _write_results_csv(path, [
        {"match_id": "M001", "team_a": "Brazil", "team_b": "Serbia",
         "actual_score_a": 2, "actual_score_b": 0},
    ])

    results = load_historical_results(path)

    assert len(results) == 1
    assert results[0].match_id == "M001"
    assert results[0].team_a == "Brazil"
    assert results[0].actual_score_a == 2
    assert results[0].actual_score_b == 0


def test_load_historical_results_applies_optional_column_defaults(tmp_path: Path) -> None:
    path = tmp_path / "results.csv"
    _write_results_csv(path, [
        {"match_id": "M001", "team_a": "A", "team_b": "B",
         "actual_score_a": 1, "actual_score_b": 1},
    ])

    results = load_historical_results(path)

    assert results[0].stage == "group stage"
    assert results[0].group == ""


def test_load_historical_results_raises_on_missing_required_column(tmp_path: Path) -> None:
    path = tmp_path / "results.csv"
    _write_results_csv(path, [
        {"match_id": "M001", "team_a": "A",
         "actual_score_a": 1, "actual_score_b": 0},  # missing team_b
    ])

    with pytest.raises(ValueError, match="team_b"):
        load_historical_results(path)


def test_load_historical_results_skips_rows_with_invalid_scores(tmp_path: Path) -> None:
    path = tmp_path / "results.csv"
    _write_results_csv(path, [
        {"match_id": "M001", "team_a": "A", "team_b": "B",
         "actual_score_a": "bad", "actual_score_b": 0},
        {"match_id": "M002", "team_a": "C", "team_b": "D",
         "actual_score_a": 1, "actual_score_b": 1},
    ])

    results = load_historical_results(path)

    assert len(results) == 1
    assert results[0].match_id == "M002"


# ---------------------------------------------------------------------------
# Unit tests: helpers
# ---------------------------------------------------------------------------

def test_parse_score_parses_valid_label() -> None:
    assert _parse_score("2-1") == (2, 1)
    assert _parse_score("0-0") == (0, 0)


def test_parse_score_returns_none_for_invalid_input() -> None:
    assert _parse_score("") is None
    assert _parse_score(None) is None
    assert _parse_score("bad") is None


def test_score_strategies_extracts_ev_default() -> None:
    row = pd.Series({
        "recommended_score": "1-0",
        "baseline_poisson_recommended_score": "1-0",
        "most_likely_scoreline": "1-0",
        "dixon_coles_recommended_score": "1-0",
        "market_consistent_recommended_score": "1-0",
        "market_consistent_status": "skipped",
        "best_expected_points": 5.5,
        "baseline_poisson_best_expected_points": 5.4,
        "most_likely_expected_points": 5.3,
        "dixon_coles_best_expected_points": 5.2,
    })

    strategies = _score_strategies(row, actual_a=1, actual_b=0)

    names = {s["strategy"] for s in strategies}
    assert "ev_default" in names
    assert "market_consistent" not in names  # MC was skipped
    ev = next(s for s in strategies if s["strategy"] == "ev_default")
    baseline = next(s for s in strategies if s["strategy"] == "baseline_poisson")
    assert ev["model_expected_points"] == pytest.approx(5.5)
    assert baseline["model_expected_points"] == pytest.approx(5.4)


def test_score_strategies_includes_market_consistent_when_mc_ran() -> None:
    row = pd.Series({
        "recommended_score": "1-0",
        "baseline_poisson_recommended_score": "1-0",
        "most_likely_scoreline": "1-0",
        "dixon_coles_recommended_score": "1-0",
        "market_consistent_recommended_score": "2-0",
        "market_consistent_status": "ok",
        "best_expected_points": 5.1,
        "baseline_poisson_best_expected_points": 5.1,
        "most_likely_expected_points": 5.0,
        "dixon_coles_best_expected_points": 5.1,
        "market_consistent_best_expected_points": 5.8,
    })

    strategies = _score_strategies(row, actual_a=2, actual_b=0)

    names = {s["strategy"] for s in strategies}
    assert "market_consistent" in names
    mc = next(s for s in strategies if s["strategy"] == "market_consistent")
    assert mc["realised_points"] == 10  # exact score
    assert mc["model_expected_points"] == pytest.approx(5.8)
    ev = next(s for s in strategies if s["strategy"] == "ev_default")
    assert ev["realised_points"] == 5  # correct result only (1-0 vs 2-0: different gd)


def test_build_summary_adds_expected_and_actual_point_totals() -> None:
    predictions = pd.DataFrame(
        [
            {
                "config": "baseline_ev",
                "strategy": "ev_default",
                "match_id": "M001",
                "realised_points": 5,
                "model_expected_points": 4.5,
                "is_exact_score": False,
                "is_correct_goal_difference": False,
                "is_correct_result": True,
            },
            {
                "config": "baseline_ev",
                "strategy": "ev_default",
                "match_id": "M002",
                "realised_points": 10,
                "model_expected_points": 6.0,
                "is_exact_score": True,
                "is_correct_goal_difference": True,
                "is_correct_result": True,
            },
        ]
    )

    summary = _build_summary(predictions)
    row = summary.iloc[0]

    assert row["matches_used"] == 2
    assert row["total_expected_points"] == pytest.approx(10.5)
    assert row["average_expected_points"] == pytest.approx(5.25)
    assert row["actual_points_sum"] == 15
    assert row["total_points"] == 15


def test_expected_vs_actual_terminal_summary_mentions_n_games() -> None:
    summary = pd.DataFrame(
        [
            {
                "config": "baseline_ev",
                "strategy": "ev_default",
                "matches_used": 3,
                "total_expected_points": 16.2,
                "average_expected_points": 5.4,
                "actual_points_sum": 17,
                "average_realised_points": 5.6667,
            }
        ]
    )

    output = _format_expected_vs_actual_summary(summary)

    assert "n=3" in output
    assert "expected total 16.20" in output
    assert "actual total 17" in output


def test_group_stage_playing_round_boundaries_are_match_number_based() -> None:
    assert _group_stage_playing_round("M001", "group stage") == "round_1"
    assert _group_stage_playing_round("M016", "group stage") == "round_1"
    assert _group_stage_playing_round("M017", "group stage") == "round_2"
    assert _group_stage_playing_round("M032", "group stage") == "round_2"
    assert _group_stage_playing_round("M033", "group stage") == "round_3"
    assert _group_stage_playing_round("M048", "group stage") == "round_3"
    assert _group_stage_playing_round("M049", "group stage") == ""
    assert _group_stage_playing_round("M016", "round of 16") == ""


def test_build_round_summary_groups_group_stage_totals() -> None:
    predictions = pd.DataFrame(
        [
            {
                "config": "baseline_ev",
                "strategy": "ev_default",
                "match_id": "M001",
                "group_stage_playing_round": "round_1",
                "realised_points": 5,
                "model_expected_points": 4.5,
                "is_exact_score": False,
                "is_correct_goal_difference": False,
                "is_correct_result": True,
            },
            {
                "config": "baseline_ev",
                "strategy": "ev_default",
                "match_id": "M016",
                "group_stage_playing_round": "round_1",
                "realised_points": 10,
                "model_expected_points": 6.0,
                "is_exact_score": True,
                "is_correct_goal_difference": True,
                "is_correct_result": True,
            },
            {
                "config": "baseline_ev",
                "strategy": "ev_default",
                "match_id": "M017",
                "group_stage_playing_round": "round_2",
                "realised_points": 1,
                "model_expected_points": 5.0,
                "is_exact_score": False,
                "is_correct_goal_difference": False,
                "is_correct_result": False,
            },
        ]
    )

    round_summary = _build_round_summary(predictions)
    round_1 = round_summary[round_summary["group_stage_playing_round"] == "round_1"].iloc[0]
    round_2 = round_summary[round_summary["group_stage_playing_round"] == "round_2"].iloc[0]
    terminal = _format_round_summary(round_summary)

    assert round_1["matches_used"] == 2
    assert round_1["round_match_range"] == "M001-M016"
    assert round_1["actual_points_sum"] == 15
    assert round_1["total_expected_points"] == pytest.approx(10.5)
    assert round_2["round_match_range"] == "M017-M032"
    assert "round_1 (M001-M016)" in terminal


# ---------------------------------------------------------------------------
# Integration tests: run_live_backtest
# ---------------------------------------------------------------------------

def test_backtest_produces_summary_and_predictions_for_single_match(
    tmp_path: Path,
) -> None:
    settings = _simple_settings(tmp_path, ["M001"])

    report = run_live_backtest(settings, export=True, progress=False)

    assert not report.predictions.empty
    assert not report.summary.empty
    assert "config" in report.predictions.columns
    assert "strategy" in report.predictions.columns
    assert "realised_points" in report.predictions.columns
    assert "model_expected_points" in report.predictions.columns
    assert "group_stage_playing_round" in report.predictions.columns
    assert "market_a_win" in report.predictions.columns
    assert "market_draw" in report.predictions.columns
    assert "lambda_a" in report.predictions.columns
    assert "ou_median_total" in report.predictions.columns
    assert "market_total_line_used" in report.predictions.columns
    assert "ev_default_score" in report.predictions.columns
    assert "most_likely_score" in report.predictions.columns
    assert "draw_vs_decisive_gap" in report.predictions.columns
    assert "draw_prone_flag" in report.predictions.columns
    assert "blowout_risk_flag" in report.predictions.columns
    assert "modal_draw_challenger_flag" in report.predictions.columns
    assert "btts_conflict_flag" in report.predictions.columns
    assert "btts_conflict_reason" in report.predictions.columns
    assert "correct_score_top_scores" in report.predictions.columns
    assert "has_other_bucket" in report.predictions.columns
    assert "correct_score_tail_mass" in report.predictions.columns
    assert "correct_score_poisson_weight" in report.predictions.columns
    assert "grid_max_goals_used" in report.predictions.columns
    ev_row = report.predictions[report.predictions["strategy"] == "ev_default"].iloc[0]
    assert ev_row["predicted_score"] == ev_row["ev_default_score"]
    assert "total_expected_points" in report.summary.columns
    assert "actual_points_sum" in report.summary.columns
    assert not report.round_summary.empty
    assert "group_stage_playing_round" in report.round_summary.columns
    assert (tmp_path / "summary.csv").exists()
    assert (tmp_path / "backtest.xlsx").exists()
    workbook = load_workbook(tmp_path / "backtest.xlsx", read_only=True)
    assert "group_stage_rounds" in workbook.sheetnames
    assert "matchday_performance" in workbook.sheetnames
    assert "ev_vs_modal_attribution" in workbook.sheetnames
    assert "draw_prone_matches" in workbook.sheetnames
    assert "blowout_risk_matches" in workbook.sheetnames
    assert "btts_conflict_matches" in workbook.sheetnames
    assert "correct_score_blend_sweep" in workbook.sheetnames
    assert "larger_grid_sweep" in workbook.sheetnames
    assert "pattern_flags_summary" in workbook.sheetnames
    assert "favourite_bucket_performance" in workbook.sheetnames
    assert "manual_review_performance" in workbook.sheetnames


def test_backtest_produces_output_for_multiple_matches(tmp_path: Path) -> None:
    settings = _simple_settings(tmp_path, ["M001", "M002", "M003"])

    report = run_live_backtest(settings, export=False, progress=False)

    assert report.predictions["match_id"].nunique() == 3
    assert not report.summary.empty


def test_backtest_records_skip_when_no_odds_file(tmp_path: Path) -> None:
    # M001 has odds; M999 has no paste file
    settings = _simple_settings(tmp_path, ["M001"])
    results_path = settings.results_path
    existing = pd.read_csv(results_path)
    extra = pd.DataFrame([{
        "match_id": "M999",
        "date": "2022-11-20",
        "stage": "group stage",
        "group": "A",
        "team_a": "X",
        "team_b": "Y",
        "actual_score_a": 1,
        "actual_score_b": 0,
    }])
    pd.concat([existing, extra]).to_csv(results_path, index=False)

    report = run_live_backtest(settings, export=False, progress=False)

    assert not report.skipped.empty
    assert "M999" in report.skipped["match_id"].values


def test_backtest_with_ah_runs_mc_challenger(tmp_path: Path) -> None:
    from wc_predictor.config import ProjectConfig

    mc_config = DEFAULT_BACKTEST_CONFIGS[3]  # mc_with_ah
    settings = _simple_settings(tmp_path, ["M001"], include_ah=True, configs=(mc_config,))

    report = run_live_backtest(settings, export=False, progress=False)

    assert not report.predictions.empty
    assert "mc_with_ah" in report.predictions["config"].values


def test_backtest_summary_ranks_configs_by_average_points(tmp_path: Path) -> None:
    settings = _simple_settings(tmp_path, ["M001", "M002"])

    report = run_live_backtest(settings, export=False, progress=False)

    assert "rank" in report.summary.columns
    assert report.summary["rank"].min() == 1


def test_classify_missing_odds_distinguishes_template_missing_and_unparseable(
    tmp_path: Path,
) -> None:
    odds = tmp_path / "odds"
    odds.mkdir()
    (odds / "M001.txt").write_text("### MATCH\nA vs B\n\n### 1X2\n<paste here>", encoding="utf-8")
    (odds / "M002.txt").write_text("### MATCH\nA vs B\n\n### 1X2\ngarbage no odds", encoding="utf-8")

    assert _classify_missing_odds(odds, "M001") == "paste_file_not_filled_in_yet"
    assert _classify_missing_odds(odds, "M002") == "no_valid_1x2_odds_parsed"
    assert _classify_missing_odds(odds, "M999") == "no_paste_file"


def test_backtest_skips_unfilled_template_with_clear_reason(tmp_path: Path) -> None:
    # M001 has real data; M002 is left as an unedited template.
    settings = _simple_settings(tmp_path, ["M001"])
    (settings.historical_odds_folder / "M002.txt").write_text(
        "### MATCH\nTeam A vs Team B\n\n### 1X2\n<paste OddsPortal 1X2 table here>",
        encoding="utf-8",
    )
    existing = pd.read_csv(settings.results_path)
    extra = pd.DataFrame([{
        "match_id": "M002", "date": "2022-11-20", "stage": "group stage",
        "group": "A", "team_a": "Team A", "team_b": "Team B",
        "actual_score_a": 1, "actual_score_b": 0,
    }])
    pd.concat([existing, extra]).to_csv(settings.results_path, index=False)

    report = run_live_backtest(settings, export=False, progress=False)

    m002_skip = report.skipped[report.skipped["match_id"] == "M002"]
    assert not m002_skip.empty
    assert m002_skip.iloc[0]["reason"] == "paste_file_not_filled_in_yet"


def test_quick_configs_are_a_subset_of_default_configs() -> None:
    default_names = {c.name for c in DEFAULT_BACKTEST_CONFIGS}
    quick_names = {c.name for c in QUICK_BACKTEST_CONFIGS}

    assert quick_names.issubset(default_names)


def test_research_configs_include_requested_sweeps() -> None:
    names = {c.name for c in RESEARCH_BACKTEST_CONFIGS}

    assert "cs_weight_1_0" in names
    assert "cs_weight_0_85" in names
    assert "cs_weight_0_75" in names
    assert "cs_weight_0_5" in names
    assert "modal_draw_fav_0_45_gap_m0_7" in names
    assert "larger_grid_15" in names
