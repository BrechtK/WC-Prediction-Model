from __future__ import annotations

from pathlib import Path
import runpy
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from wc_predictor.dixon_coles_rho_comparison import (
    DEFAULT_DIXON_COLES_RHOS,
    DETAIL_COLUMNS,
    _build_rho_summary,
    compare_dixon_coles_rhos,
    export_dixon_coles_rho_comparison,
    format_dixon_coles_rho_comparison_summary,
    parse_rho_grid,
)
from wc_predictor.market_data import load_odds
from wc_predictor.workflow import run_prediction_workflow

EXAMPLES = Path("data/examples")
CORE_FIXTURES = Path("tests/fixtures/oddsportal_core_pastes")
RUN_SCRIPT = Path("scripts/compare_dixon_coles_rho.py").resolve()
TEST_RHOS = (-0.10, 0.00, 0.10)


def _write_metadata(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(
        [
            {
                "match_id": "M001",
                "date": "2026-06-12",
                "stage": "group",
                "group": "A",
                "team_a": "Alpha",
                "team_b": "Beta",
                "bookmaker": "bet365",
            }
        ]
    ).to_excel(path, index=False)


def _write_core_pastes(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    fixtures = {
        "1x2": CORE_FIXTURES / "MCORE_1x2.txt",
        "over_under": CORE_FIXTURES / "MCORE_over_under.txt",
        "btts": CORE_FIXTURES / "MCORE_btts.txt",
    }
    for market, source in fixtures.items():
        (folder / f"M001_{market}.txt").write_text(source.read_text(encoding="utf-8"), encoding="utf-8")


def test_default_and_custom_rho_grids() -> None:
    assert DEFAULT_DIXON_COLES_RHOS == (-0.20, -0.15, -0.10, -0.05, 0.00, 0.05, 0.10, 0.15, 0.20)
    assert parse_rho_grid("0.10, -0.10, 0, 0.10") == (-0.10, 0.0, 0.10)


def test_rho_zero_reproduces_poisson_and_nonzero_rho_changes_low_scores() -> None:
    odds = load_odds(EXAMPLES / "example_odds.csv")
    baseline = run_prediction_workflow(odds)
    comparison = compare_dixon_coles_rhos(EXAMPLES / "example_odds.csv", rhos=TEST_RHOS)

    for match_id, group in comparison.details.groupby("match_id"):
        zero = group[np.isclose(group["rho"], 0.0)].iloc[0]
        negative = group[np.isclose(group["rho"], -0.10)].iloc[0]
        matrix = baseline.baseline_score_matrices[str(match_id)]
        assert zero["dixon_coles_recommended_score"] == zero["baseline_poisson_recommended_score"]
        assert zero["p_dc_0_0"] == matrix.exact_score_probability(0, 0)
        assert zero["p_dc_1_0"] == matrix.exact_score_probability(1, 0)
        assert zero["p_dc_0_1"] == matrix.exact_score_probability(0, 1)
        assert zero["p_dc_1_1"] == matrix.exact_score_probability(1, 1)
        assert negative["p_dc_0_0"] != zero["p_dc_0_0"]
        assert negative["p_dc_1_0"] != zero["p_dc_1_0"]
        assert negative["p_dc_0_1"] != zero["p_dc_0_1"]
        assert negative["p_dc_1_1"] != zero["p_dc_1_1"]
        assert group["dixon_coles_matrix_normalised"].all()


def test_rho_summary_detects_recommendation_changes_and_stable_interval() -> None:
    details = pd.DataFrame(
        [
            {
                "match_id": "M001",
                "team_a": "A",
                "team_b": "B",
                "rho": rho,
                "dixon_coles_recommended_score": "1-0" if -0.05 <= rho <= 0.15 else "1-1",
                "final_live_recommended_score": "1-0",
                "dixon_coles_matrix_normalised": True,
            }
            for rho in DEFAULT_DIXON_COLES_RHOS
        ]
    )

    summary = _build_rho_summary(details).iloc[0]

    assert summary["recommendation_at_rho_0"] == "1-0"
    assert summary["recommendation_changes_across_rho"] == "yes"
    assert summary["first_rho_where_recommendation_changes_from_rho_0"] == -0.10
    assert summary["stable_rho_interval_around_0"] == "[-0.05, 0.15]"
    assert not summary["live_recommendation_agrees_with_all_dixon_coles"]
    assert summary["live_recommendation_agrees_with_most_dixon_coles"]


def test_rho_comparison_exports_workbook_and_console_summary(tmp_path: Path) -> None:
    comparison = compare_dixon_coles_rhos(EXAMPLES / "example_odds.csv", rhos=TEST_RHOS)

    output_path = export_dixon_coles_rho_comparison(comparison, tmp_path / "comparison.xlsx")
    workbook = load_workbook(output_path)

    assert output_path.exists()
    assert workbook.sheetnames == ["details", "summary"]
    assert workbook["details"].freeze_panes == "A2"
    assert [cell.value for cell in workbook["details"][1]] == DETAIL_COLUMNS
    text = format_dixon_coles_rho_comparison_summary(comparison, output_path)
    assert "Dixon-Coles challenger rho sensitivity" in text
    assert "Matches compared: 3" in text
    assert f"Output workbook path: {output_path}" in text


def test_rho_comparison_script_uses_live_pastes_and_writes_default_workbook(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    _write_metadata(tmp_path / "input/prepared/world_cup_odds.xlsx")
    _write_core_pastes(tmp_path / "cache/split_pastes")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys, "argv", ["compare_dixon_coles_rho.py"])

    runpy.run_path(str(RUN_SCRIPT), run_name="__main__")

    output = capsys.readouterr().out
    assert "Dixon-Coles challenger rho sensitivity" in output
    assert "Matches compared: 1" in output
    assert (tmp_path / "output/dixon_coles_rho_comparison.xlsx").exists()
