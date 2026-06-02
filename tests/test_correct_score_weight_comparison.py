from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from wc_predictor.correct_score_weight_comparison import (
    DEFAULT_CORRECT_SCORE_WEIGHTS,
    DETAIL_COLUMNS,
    _build_sensitivity_summary,
    compare_correct_score_weights,
    export_correct_score_weight_comparison,
    format_correct_score_weight_comparison_summary,
    parse_weight_grid,
)


REFERENCE_WEIGHTS = (1.00, 0.85, 0.75, 0.50, 0.00)


def _write_dummy_inputs(tmp_path: Path) -> tuple[Path, Path]:
    odds_path = tmp_path / "world_cup_odds.xlsx"
    correct_score_path = tmp_path / "world_cup_correct_score_odds.csv"
    pd.DataFrame(
        [
            {
                "match_id": "STABLE",
                "date": "2026-06-12",
                "stage": "group",
                "group": "A",
                "team_a": "Stable A",
                "team_b": "Stable B",
                "bookmaker": "Example",
                "odds_a_win": 1.65,
                "odds_draw": 3.90,
                "odds_b_win": 5.50,
            },
            {
                "match_id": "CHANGING",
                "date": "2026-06-13",
                "stage": "group",
                "group": "A",
                "team_a": "Changing A",
                "team_b": "Changing B",
                "bookmaker": "Example",
                "odds_a_win": 1.65,
                "odds_draw": 3.90,
                "odds_b_win": 5.50,
            },
        ]
    ).to_excel(odds_path, index=False)
    scorelines = [
        (0, 0),
        (1, 0),
        (0, 1),
        (1, 1),
        (2, 0),
        (0, 2),
        (2, 1),
        (1, 2),
        (3, 0),
        (0, 3),
    ]
    stable_odds = [15.0, 1.60, 30.0, 12.0, 10.0, 50.0, 20.0, 50.0, 40.0, 100.0]
    changing_odds = [30.0, 30.0, 50.0, 30.0, 1.40, 100.0, 30.0, 100.0, 40.0, 100.0]
    rows = []
    for match_id, odds in [("STABLE", stable_odds), ("CHANGING", changing_odds)]:
        rows.extend(
            {
                "match_id": match_id,
                "bookmaker": "Example",
                "score_a": score_a,
                "score_b": score_b,
                "decimal_odds": decimal_odds,
            }
            for (score_a, score_b), decimal_odds in zip(scorelines, odds, strict=True)
        )
    pd.DataFrame(rows).to_csv(correct_score_path, index=False)
    return odds_path, correct_score_path


def test_default_and_custom_weight_grids() -> None:
    assert DEFAULT_CORRECT_SCORE_WEIGHTS == tuple(round(value / 100, 2) for value in range(100, -1, -5))
    assert parse_weight_grid("0, 0.85, 1, 0.85") == (1.0, 0.85, 0.0)


def test_stable_range_and_first_change_use_the_discrete_weight_path() -> None:
    weights = [1.00, 0.95, 0.90, 0.85, 0.80, 0.75, 0.70, 0.65, 0.60, 0.55]
    details = pd.DataFrame(
        [
            {
                "match_id": "M001",
                "team_a": "A",
                "team_b": "B",
                "weight": weight,
                "recommended_score": "1-0" if weight >= 0.65 else "2-0",
                "recommended_qualifier": "",
                "ev_gap_best_vs_second": 0.1,
            }
            for weight in weights
        ]
    )

    summary = _build_sensitivity_summary(details).iloc[0]

    assert summary["first_weight_where_recommendation_changes_from_poisson"] == 0.60
    assert summary["stable_weight_range_around_0_85"] == "[0.65, 1.00]"


def test_weight_comparison_finds_stable_and_changing_matches_and_exports_workbook(tmp_path: Path) -> None:
    odds_path, correct_score_path = _write_dummy_inputs(tmp_path)

    comparison = compare_correct_score_weights(
        odds_path,
        correct_score_path,
        weights=REFERENCE_WEIGHTS,
    )
    summary = comparison.sensitivity.set_index("match_id")

    assert summary.loc["STABLE", "recommendation_changes"] == "no"
    assert summary.loc["CHANGING", "recommendation_changes"] == "yes"
    assert summary.loc["CHANGING", "recommendation_at_w_1_00"] == "1-0"
    assert summary.loc["CHANGING", "recommendation_at_w_0_00"] == "2-0"
    assert comparison.details["weight"].nunique() == len(REFERENCE_WEIGHTS)

    output_path = export_correct_score_weight_comparison(comparison, tmp_path / "comparison.xlsx")
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["details", "sensitivity"]
    assert workbook["details"].freeze_panes == "A2"
    assert [cell.value for cell in workbook["details"][1]] == DETAIL_COLUMNS

    text = format_correct_score_weight_comparison_summary(comparison, output_path)
    assert "Matches compared: 2" in text
    assert "Matches where recommendation changes across weights: 1" in text
    assert "Matches stable for all weights: 1" in text


def test_weight_comparison_script_runs_with_custom_grid(tmp_path: Path) -> None:
    odds_path, correct_score_path = _write_dummy_inputs(tmp_path)
    output_path = tmp_path / "comparison.xlsx"

    completed = subprocess.run(
        [
            sys.executable,
            "scripts/compare_correct_score_weights.py",
            "--odds",
            str(odds_path),
            "--correct-score-odds",
            str(correct_score_path),
            "--weights",
            "1,0.85,0.75,0.5,0",
            "--output",
            str(output_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )

    assert "Correct-score blend-weight sensitivity" in completed.stdout
    assert output_path.exists()
